import re
import time
import unicodedata
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple
from urllib.parse import quote
import logging
import sys
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
from playwright.sync_api import sync_playwright

#Constantes globais de acesso
LOGIN_URL = "https://servicosti.ebserh.gov.br/#/login"
USER_URL_TEMPLATE = "https://servicosti.ebserh.gov.br/#/usuarios/{search_value}"

#Constantes globais de validação de acesso (login)
LOGIN_BUTTON_TEXT = "Entrar"
LOGIN_PASSWORD_SELECTOR = 'input[type="password"]'
PAGE_READY_TIMEOUT_MS = 15000
LOGIN_COMPLETION_TIMEOUT_MS = 10000
LOGIN_COMPLETION_POLL_INTERVAL_MS = 150
LOGIN_FAILED_WARNING_SELECTOR = (
    "[role='alert'], "
    ".alert, "
    ".alert-danger, "
    ".alert-warning, "
    ".toast, "
    ".toast-error, "
    ".toast-message, "
    ".swal2-popup"
)
LOGIN_FAILED_WARNING_PATTERN = re.compile(
    r"("
    r"login|entrar|autentic|credenc|usuario|usuário|senha"
    r").{0,80}("
    r"falh|inval|invál|incorret|negad|bloquead|expirad"
    r")|("
    r"falh|inval|invál|incorret|negad|bloquead|expirad"
    r").{0,80}("
    r"login|entrar|autentic|credenc|usuario|usuário|senha"
    r")",
    re.IGNORECASE,
)

# Constantes globais da tela de usuário
USER_UPDATE_BUTTON_TEXT = "Atualizar dados"

USER_PAGE_READY_TIMEOUT_MS = 15000
USER_DATE_INPUT_TIMEOUT_MS = 7000
USER_UPDATE_BUTTON_TIMEOUT_MS = 7000
USER_SAVE_RESPONSE_TIMEOUT_MS = 10000
USER_SAVE_CONFIRMATION_TIMEOUT_MS = 3000
USER_SAVE_REQUEST_METHODS = {
    "POST",
    "PUT",
    "PATCH",
}

USER_NOT_FOUND_SELECTOR = (
    "[role='alert'], "
    ".alert, "
    ".alert-danger, "
    ".alert-warning, "
    ".toast, "
    ".toast-error, "
    ".toast-message, "
    ".well, "
    ".swal2-popup"
)
USER_NOT_FOUND_PATTERN = re.compile(
    r"("
    r"usu[aá]rio|user"
    r").{0,80}("
    r"n[aã]o encontrado|não encontrado|nao encontrado|inexistente"
    r")|("
    r"n[aã]o encontrado|não encontrado|nao encontrado|inexistente"
    r").{0,80}("
    r"usu[aá]rio|user"
    r")",
    re.IGNORECASE,
)

USER_SAVE_CONFIRMATION_SELECTOR = (
    "[role='alert'], "
    ".alert, "
    ".alert-success, "
    ".alert-info, "
    ".toast, "
    ".toast-success, "
    ".toast-message, "
    ".swal2-popup"
)

USER_SAVE_CONFIRMATION_PATTERN = re.compile(
    r"("
    r"salv|atualiz|alterad|sucesso|prorrog|dados"
    r")",
    re.IGNORECASE,
)

#Constantes globais de data
DATE_INPUT_NAME = "__/__/____"
EXPIRATION_DATE_FORMAT = "%d/%m/%Y"
EXPIRATION_DATE_PATTERN = re.compile(r"^\d{2}/\d{2}/\d{4}$")

#Constantes globais de relatórios
REPORT_COLUMN_USER = "usuário"
REPORT_COLUMN_NAME = "relatório"
REPORT_HEADERS = (REPORT_COLUMN_USER,REPORT_COLUMN_NAME,)
USER_COLUMN_CANDIDATES = ("usuário","usuario","login","rede",)

#Constantes globais de status de lote
USER_NOT_FOUND_MESSAGE = "Usuário não Encontrado"
ERROR_MESSAGE_PREFIX = "Erro:"

STATUS_SUCESSO = "sucesso"
STATUS_NAO_ENCONTRADO = "nao_encontrado"
STATUS_ERRO = "erro"

#Constantes globais de usuários
USER_LOGIN_ALLOWED_PATTERN = re.compile(r"^[A-Za-z0-9._@'-]+$")
USER_URL_SAFE_CHARS = "._@-'"
BATCH_ENTRY_PREPARED_USER = "_prepared_user"

#Constantes globais de planilha
SUPPORTED_EXTENSIONS = (".xlsx",)
Row = Dict[str, Any]

# Constantes globais de layout XLSX
XLSX_HEADER_ROW = 1
XLSX_FREEZE_PANES_CELL = "A2"
XLSX_MIN_COLUMN_WIDTH = 12
XLSX_MAX_COLUMN_WIDTH = 60
XLSX_COLUMN_PADDING = 2

# Constantes globais de logging
LOGGER_NAME = "Prorrogador"
LOG_FORMAT = "%(asctime)s [%(levelname)s] %(message)s"
LOG_DATE_FORMAT = "%d/%m/%Y %H:%M:%S"

logger = logging.getLogger(LOGGER_NAME)
logger.addHandler(logging.NullHandler())


# -----------------------------
# Logging
# -----------------------------

# Obtém stream disponível para logging em terminal.
def get_stream_for_logging() -> Optional[Any]:
    return (
        sys.stdout
        or sys.stderr
        or sys.__stdout__
        or sys.__stderr__
    )


# Configura logging da automação.
def configure_automation_logging(
    show_terminal_logs: bool = False
) -> logging.Logger:
    automation_logger = logging.getLogger(LOGGER_NAME)
    automation_logger.setLevel(logging.INFO)
    automation_logger.propagate = False

    for handler in list(automation_logger.handlers):
        automation_logger.removeHandler(handler)
        handler.close()

    if show_terminal_logs:
        formatter = logging.Formatter(
            LOG_FORMAT,
            datefmt=LOG_DATE_FORMAT
        )

        stream = get_stream_for_logging()

        if stream is not None:
            stream_handler = logging.StreamHandler(stream)
            stream_handler.setLevel(logging.INFO)
            stream_handler.setFormatter(formatter)
            automation_logger.addHandler(stream_handler)

    else:
        automation_logger.addHandler(logging.NullHandler())

    automation_logger.info(
        "Logging configurado. Terminal: %s.",
        "sim" if show_terminal_logs else "não"
    )

    return automation_logger

# -----------------------------
# Planilha / relatório
# -----------------------------

# Normaliza nome de coluna.
def normalize_column_name(value: Any) -> str:
    text = "" if value is None else str(value).strip()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(
        char for char in text
        if not unicodedata.combining(char)
    )
    return text.lower()


# Identifica coluna de usuários.
def identify_user_column(headers: Sequence[str]) -> str:
    normalized_candidates = {
        normalize_column_name(candidate)
        for candidate in USER_COLUMN_CANDIDATES
    }

    for header in headers:
        if normalize_column_name(header) in normalized_candidates:
            return header

    expected_columns = ", ".join(USER_COLUMN_CANDIDATES)
    raise ValueError(
        "Coluna de usuários não encontrada. "
        f"A planilha deve conter uma destas colunas: {expected_columns}."
    )


# Gera nome do relatório.
def generate_report_filename(
    extension: str,
    now: Optional[datetime] = None
) -> str:
    reference_date = now or datetime.now()
    clean_extension = extension.lower().strip()

    if not clean_extension.startswith("."):
        clean_extension = f".{clean_extension}"

    timestamp = reference_date.strftime("%d_%m_%y_%Hh%Mm%S")

    return f"Resultado_{timestamp}{clean_extension}"


# Constrói caminho do relatório.
def build_report_path(
    report_directory: str,
    source_spreadsheet_path: str,
    now: Optional[datetime] = None
) -> Path:
    report_dir = Path(report_directory).expanduser()

    if not report_dir.exists():
        raise FileNotFoundError(
            f"A pasta de relatório não existe: {report_dir}"
        )

    if not report_dir.is_dir():
        raise NotADirectoryError(
            f"O caminho de relatório não é uma pasta: {report_dir}"
        )

    extension = validate_spreadsheet_extension(source_spreadsheet_path)
    filename = generate_report_filename(extension, now)
    report_path = report_dir / filename

    return get_available_report_path(report_path)


# Obtém caminho disponível.
def get_available_report_path(report_path: Path) -> Path:
    if not report_path.exists():
        return report_path

    counter = 2
    stem = report_path.stem
    suffix = report_path.suffix
    parent = report_path.parent

    while True:
        candidate = parent / f"{stem}_{counter}{suffix}"

        if not candidate.exists():
            return candidate

        counter += 1


# Valida extensão da planilha.
def validate_spreadsheet_extension(spreadsheet_path: str) -> str:
    extension = Path(spreadsheet_path).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(SUPPORTED_EXTENSIONS)
        raise ValueError(
            f"Formato de planilha não suportado: {extension}. "
            f"Use apenas: {supported}."
        )

    return extension


# Lê dados da planilha.
def read_spreadsheet(spreadsheet_path: str) -> Tuple[List[str], List[Row]]:
    source_path = Path(spreadsheet_path).expanduser()

    if not source_path.exists():
        raise FileNotFoundError(
            f"Planilha não encontrada: {source_path}"
        )

    validate_spreadsheet_extension(str(source_path))

    return read_xlsx(source_path)


# Lê arquivo Excel.
def read_xlsx(source_path: Path) -> Tuple[List[str], List[Row]]:
    workbook = load_workbook(source_path, data_only=True)
    worksheet = workbook.active

    if worksheet is None:
        raise ValueError("A planilha não possui aba ativa.")

    raw_headers = [
        cell.value
        for cell in next(worksheet.iter_rows(min_row=1, max_row=1))
    ]
    headers = make_unique_headers(raw_headers)
    rows: List[Row] = []

    for raw_row in worksheet.iter_rows(min_row=2, values_only=True):
        if is_empty_row(raw_row):
            continue

        row = build_row(headers, raw_row)
        rows.append(row)

    return headers, rows

# Garante cabeçalhos únicos.
def make_unique_headers(raw_headers: Sequence[Any]) -> List[str]:
    if not raw_headers:
        raise ValueError("A planilha não possui cabeçalho.")

    headers: List[str] = []
    seen: Dict[str, int] = {}

    for index, raw_header in enumerate(raw_headers, start=1):
        header = "" if raw_header is None else str(raw_header).strip()

        if not header:
            header = f"coluna_{index}"

        if header in seen:
            seen[header] += 1
            header = f"{header}_{seen[header]}"
        else:
            seen[header] = 1

        headers.append(header)

    return headers


# Constrói linha de dados.
def build_row(headers: Sequence[str], raw_row: Sequence[Any]) -> Row:
    row: Row = {}

    for index, header in enumerate(headers):
        row[header] = raw_row[index] if index < len(raw_row) else ""

    return row


# Verifica linha vazia.
def is_empty_row(raw_row: Sequence[Any]) -> bool:
    return all(
        value is None or str(value).strip() == ""
        for value in raw_row
    )


# Salva relatório em arquivo.
def write_report(
    source_spreadsheet_path: str,
    report_path: Path,
    rows: Sequence[Row]
) -> None:
    validate_spreadsheet_extension(source_spreadsheet_path)
    report_headers = get_report_headers()

    write_xlsx_report(report_path, report_headers, rows)

# Obtém cabeçalhos do relatório.
def get_report_headers() -> List[str]:
    return list(REPORT_HEADERS)

# Constrói linha limpa do relatório.
def build_report_row(user: Any, report: Any) -> Row:
    clean_user = "" if user is None else str(user).strip()
    clean_report = "" if report is None else str(report).strip()

    return {
        REPORT_COLUMN_USER: clean_user,
        REPORT_COLUMN_NAME: clean_report,
    }


# Salva relatório em Excel.
def write_xlsx_report(
    report_path: Path,
    headers: Sequence[str],
    rows: Sequence[Row]
) -> None:
    workbook = Workbook()
    worksheet = workbook.active

    if worksheet is None:
        raise ValueError("A planilha não possui aba ativa.")

    worksheet.title = "Relatório"

    worksheet.append(list(headers))

    for row in rows:
        worksheet.append([
            row.get(header, "")
            for header in headers
        ])

    apply_xlsx_report_layout(worksheet)
    workbook.save(report_path)

# Aplica filtros, congelamento e largura automática no XLSX.
def apply_xlsx_report_layout(worksheet: Worksheet) -> None:
    if worksheet.max_row < XLSX_HEADER_ROW or worksheet.max_column < 1:
        return

    worksheet.freeze_panes = XLSX_FREEZE_PANES_CELL
    worksheet.auto_filter.ref = build_xlsx_filter_range(worksheet)
    autofit_xlsx_columns(worksheet)


# Monta intervalo de filtros do relatório.
def build_xlsx_filter_range(worksheet: Worksheet) -> str:
    last_column = get_column_letter(worksheet.max_column)
    return f"A{XLSX_HEADER_ROW}:{last_column}{worksheet.max_row}"


# Ajusta largura das colunas conforme o maior conteúdo.
def autofit_xlsx_columns(worksheet: Worksheet) -> None:
    for column_cells in worksheet.columns:
        if not column_cells:
            continue
        col_idx = column_cells[0].column
        if not isinstance(col_idx, int):
            continue
        column_letter = get_column_letter(col_idx)
        max_length = max(
            get_xlsx_cell_text_length(cell.value)
            for cell in column_cells
        )
        width = max_length + XLSX_COLUMN_PADDING
        worksheet.column_dimensions[column_letter].width = min(
            max(width, XLSX_MIN_COLUMN_WIDTH),
            XLSX_MAX_COLUMN_WIDTH,
        )

# Calcula tamanho visível de uma célula.
def get_xlsx_cell_text_length(value: Any) -> int:
    if value is None:
        return 0

    lines = str(value).splitlines() or [""]
    return max(len(line) for line in lines)

# -----------------------------
# Login alvo / validação
# -----------------------------

# Normaliza valor de usuário vindo da UI ou planilha.
def normalize_user_value(value: Any) -> str:
    if value is None:
        return ""

    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


# Prepara usuário para execução individual.
def prepare_user_value(user_value: Any) -> str:
    if isinstance(user_value, bool):
        raise ValueError(
            "Usuário inválido: valor booleano não é aceito."
        )

    user = normalize_user_value(user_value)

    if not user:
        raise ValueError("Usuário não informado.")

    if any(char.isspace() for char in user):
        raise ValueError(
            "Usuário inválido: não use espaços."
        )

    if not USER_LOGIN_ALLOWED_PATTERN.fullmatch(user):
        raise ValueError(
            "Usuário inválido: use apenas letras, números, ponto, "
            "hífen, sublinhado, @ ou apóstrofo."
        )

    return user

# Monta URL segura para usuário já validado/preparado.
def build_user_url(prepared_user_value: str) -> str:
    encoded_user = quote(
        prepared_user_value,
        safe=USER_URL_SAFE_CHARS
    )

    return USER_URL_TEMPLATE.format(
        search_value=encoded_user
    )

# -----------------------------
# Data / validação
# -----------------------------

# Normaliza e valida a nova data de expiração.
def normalize_expiration_date(value: Any) -> str:
    if value is None:
        raise ValueError("Nova data não informada.")

    if isinstance(value, bool):
        raise ValueError("Nova data inválida: valor booleano não é aceito.")

    if isinstance(value, datetime):
        return value.strftime(EXPIRATION_DATE_FORMAT)

    if isinstance(value, date):
        return value.strftime(EXPIRATION_DATE_FORMAT)

    text = str(value).strip()

    if not text:
        raise ValueError("Nova data não informada.")

    digits = "".join(
        char for char in text
        if char.isdigit()
    )

    if text.isdigit() and len(text) == 8:
        text = f"{digits[:2]}/{digits[2:4]}/{digits[4:]}"

    if not EXPIRATION_DATE_PATTERN.fullmatch(text):
        raise ValueError(
            "Nova data inválida: use o formato dd/mm/aaaa."
        )

    try:
        parsed_date = datetime.strptime(
            text,
            EXPIRATION_DATE_FORMAT
        )
    except ValueError as exc:
        raise ValueError(
            "Nova data inválida: informe uma data real no formato dd/mm/aaaa."
        ) from exc

    return parsed_date.strftime(EXPIRATION_DATE_FORMAT)   

# -----------------------------
# Autenticação
# -----------------------------

# Obtém locator do campo de login.
def get_login_textbox_locator(page):
    return page.get_by_role("textbox").first


# Obtém locator do campo de senha do login.
def get_login_password_locator(page):
    return page.locator(LOGIN_PASSWORD_SELECTOR)


# Aguarda a tela de login carregar após DOMContentLoaded.
def wait_for_login_page_ready(page) -> None:
    get_login_textbox_locator(page).wait_for(
        state="visible",
        timeout=PAGE_READY_TIMEOUT_MS
    )
    get_login_password_locator(page).wait_for(
        state="visible",
        timeout=PAGE_READY_TIMEOUT_MS
    )


# Obtém locator do aviso de falha no login.
def get_login_failed_warning_locator(page):
    return page.locator(LOGIN_FAILED_WARNING_SELECTOR).filter(
        has_text=LOGIN_FAILED_WARNING_PATTERN
    )


# Extrai texto visível do aviso de falha no login.
def get_visible_login_failed_warning_text(page) -> Optional[str]:
    warnings = get_login_failed_warning_locator(page)

    for index in range(warnings.count()):
        warning = warnings.nth(index)

        if warning.is_visible():
            return " ".join(
                warning.inner_text(timeout=500).split()
            )

    return None


# Aguarda o login concluir sem usar espera fixa.
def wait_for_login_completion(page) -> None:
    deadline = time.monotonic() + (
        LOGIN_COMPLETION_TIMEOUT_MS / 1000
    )

    while time.monotonic() < deadline:
        failed_warning_text = get_visible_login_failed_warning_text(page)

        if failed_warning_text:
            raise RuntimeError(
                "Login falhou: aviso de falha detectado. "
                f"Mensagem do sistema: {failed_warning_text}"
            )

        if not get_login_password_locator(page).is_visible():
            return

        page.wait_for_timeout(LOGIN_COMPLETION_POLL_INTERVAL_MS)

    raise RuntimeError(
        "Login não concluído no tempo esperado. "
        "Verifique as credenciais, a disponibilidade do sistema "
        "ou mudança no fluxo de autenticação."
    )


# Realiza login no sistema.
def login_to_system(page, login: str, password: str) -> None:
    logger.info("Acessando tela de login.")

    page.goto(
        LOGIN_URL,
        wait_until="domcontentloaded"
    )

    wait_for_login_page_ready(page)

    logger.info("Tela de login carregada. Preenchendo credenciais.")

    get_login_textbox_locator(page).fill(login)
    get_login_password_locator(page).fill(password)

    page.get_by_role(
        "button",
        name=LOGIN_BUTTON_TEXT
    ).click()

    logger.info("Login submetido. Aguardando conclusão.")

    wait_for_login_completion(page)


# -----------------------------
# Tela de usuário / navegação / seletores / esperas observáveis
# -----------------------------

# Abre a tela de prorrogação do usuário.
def open_user_page(page, prepared_user_value: str) -> None:
    page.goto(
        build_user_url(prepared_user_value),
        wait_until="domcontentloaded"
    )


# Obtém locator do campo de data de expiração.
def get_expiration_date_locator(page):
    return page.get_by_role(
        "textbox",
        name=DATE_INPUT_NAME
    ).first


# Obtém locator do botão de atualização.
def get_update_button_locator(page):
    return page.get_by_role(
        "button",
        name=USER_UPDATE_BUTTON_TEXT
    )


# Obtém locator de mensagem explícita de usuário não encontrado.
def get_user_not_found_locator(page):
    return page.locator(USER_NOT_FOUND_SELECTOR).filter(
        has_text=USER_NOT_FOUND_PATTERN
    )


# Obtém locator de confirmação pós-salvamento.
def get_user_save_confirmation_locator(page):
    return page.locator(USER_SAVE_CONFIRMATION_SELECTOR).filter(
        has_text=USER_SAVE_CONFIRMATION_PATTERN
    )


# Aguarda a tela do usuário ficar pronta sem usar espera fixa.
def wait_for_user_page_ready(page) -> bool:
    expiration_date = get_expiration_date_locator(page)
    not_found = get_user_not_found_locator(page)

    try:
        expiration_date.or_(not_found).first.wait_for(
            state="visible",
            timeout=USER_PAGE_READY_TIMEOUT_MS
        )
    except PlaywrightTimeoutError:
        return False

    if not expiration_date.is_visible():
        return False

    try:
        get_update_button_locator(page).wait_for(
            state="visible",
            timeout=USER_UPDATE_BUTTON_TIMEOUT_MS
        )
    except PlaywrightTimeoutError:
        return False

    return True


# Preenche a nova data de expiração.
def fill_expiration_date(page, expiration_date: str) -> None:
    expiration_date_input = get_expiration_date_locator(page)

    expiration_date_input.wait_for(
        state="visible",
        timeout=USER_DATE_INPUT_TIMEOUT_MS
    )
    expiration_date_input.click()
    page.keyboard.press("Escape")
    expiration_date_input.press("ControlOrMeta+A")
    expiration_date_input.fill("")
    expiration_date_input.type(expiration_date)


# Identifica resposta HTTP provável de salvamento.
def is_user_save_response(response) -> bool:
    request = response.request

    if request.method not in USER_SAVE_REQUEST_METHODS:
        return False

    return 200 <= response.status < 500


# Aguarda confirmação visual pós-salvamento, quando existir.
def wait_for_user_save_confirmation_if_available(page) -> Optional[str]:
    confirmation = get_user_save_confirmation_locator(page).first

    try:
        confirmation.wait_for(
            state="visible",
            timeout=USER_SAVE_CONFIRMATION_TIMEOUT_MS
        )
    except PlaywrightTimeoutError:
        return None

    return " ".join(
        confirmation.inner_text(timeout=500).split()
    )

# Clica em atualizar e aguarda estado observável de pós-salvamento.
def click_update_button(page) -> None:
    update_button = get_update_button_locator(page)

    logger.info("Acionando botão de atualização.")

    try:
        with page.expect_response(
            is_user_save_response,
            timeout=USER_SAVE_RESPONSE_TIMEOUT_MS
        ) as response_info:
            update_button.click(timeout=USER_UPDATE_BUTTON_TIMEOUT_MS)

        response = response_info.value

        logger.info(
            "Resposta de salvamento recebida. HTTP %s.",
            response.status
        )

        if response.status >= 400:
            raise RuntimeError(
                "Falha ao salvar a prorrogação. "
                f"Resposta HTTP: {response.status}."
            )

        confirmation_text = wait_for_user_save_confirmation_if_available(page)

        if confirmation_text:
            logger.info(
                "Confirmação visual detectada: %s.",
                confirmation_text
            )

        return

    except PlaywrightTimeoutError:
        logger.warning(
            "Nenhuma resposta HTTP de salvamento foi detectada no prazo. "
            "Verificando confirmação visual."
        )

        confirmation_text = wait_for_user_save_confirmation_if_available(page)

        if confirmation_text:
            logger.info(
                "Confirmação visual detectada: %s.",
                confirmation_text
            )
            return

        try:
            update_button.wait_for(
                state="visible",
                timeout=USER_UPDATE_BUTTON_TIMEOUT_MS
            )
        except PlaywrightTimeoutError as exc:
            raise RuntimeError(
                "Não foi possível confirmar o salvamento da prorrogação. "
                "O botão de atualização não voltou a ficar disponível."
            ) from exc

        raise RuntimeError(
            "Não foi possível confirmar o salvamento da prorrogação. "
            "Nenhuma resposta HTTP ou confirmação visual foi detectada."
        )

# Prorroga data usando usuário já validado/preparado.
def process_user_prepared_value(
    page,
    prepared_user_value: str,
    prepared_expiration_date: str
) -> str:
    user = "" if prepared_user_value is None else str(
        prepared_user_value
    ).strip()

    if not user:
        raise ValueError("Usuário preparado não informado.")

    expiration_date = "" if prepared_expiration_date is None else str(
        prepared_expiration_date
    ).strip()

    if not expiration_date:
        raise ValueError("Nova data preparada não informada.")

    logger.info("Abrindo tela do usuário: %s.", user)

    open_user_page(page, user)

    if not wait_for_user_page_ready(page):
        logger.warning("Usuário não encontrado: %s.", user)
        return USER_NOT_FOUND_MESSAGE

    logger.info(
        "Tela do usuário carregada. Preenchendo nova data: %s.",
        expiration_date
    )

    fill_expiration_date(page, expiration_date)
    click_update_button(page)

    logger.info(
        "Usuário %s prorrogado com sucesso para %s.",
        user,
        expiration_date
    )

    return f"Data prorrogada para {expiration_date}"

# -----------------------------
# Resumo de lote
# -----------------------------

# Classifica a mensagem de relatório de uma linha em uma categoria de status.
def classify_batch_row_status(report_message: str) -> str:
    if report_message.startswith(ERROR_MESSAGE_PREFIX):
        return STATUS_ERRO

    if report_message == USER_NOT_FOUND_MESSAGE:
        return STATUS_NAO_ENCONTRADO

    return STATUS_SUCESSO


# Resume mensagens de resultado por status.
def summarize_messages(messages: Sequence[str]) -> str:
    contagem = Counter(
        classify_batch_row_status(message)
        for message in messages
    )
    total = len(messages)

    return (
        f"Total: {total}. "
        f"Sucesso: {contagem[STATUS_SUCESSO]}. "
        f"Não encontrado: {contagem[STATUS_NAO_ENCONTRADO]}. "
        f"Erro: {contagem[STATUS_ERRO]}."
    )


# Resume as linhas do lote por status.
def summarize_batch_results(batch_rows: Sequence[Row]) -> str:
    messages = [
        row.get(REPORT_COLUMN_NAME, "")
        for row in batch_rows
    ]

    return summarize_messages(messages)


# Define a cor de status conforme uma ou mais mensagens (unitário ou lote).
def result_color(messages: Sequence[str]) -> str:
    contagem = Counter(
        classify_batch_row_status(message)
        for message in messages
    )

    if contagem[STATUS_ERRO]:
        return "red"

    if contagem[STATUS_NAO_ENCONTRADO]:
        return "orange"

    return "green"


# Formata o detalhe de cada usuário da execução unitária multi-usuário.
def format_multi_results(
    users: Sequence[str],
    messages: Sequence[str]
) -> str:
    blocos = [
        f"{index}. Usuário: {user}\n{message}"
        for index, (user, message) in enumerate(zip(users, messages), start=1)
    ]

    return "\n\n".join(blocos)


# -----------------------------
# Entry points
# -----------------------------

# Executa automação individual.
def run_automation(
    login: str,
    password: str,
    search_value: str,
    expiration_date: str,
    show_terminal_logs: bool = False,
    mostrar_browser: bool = True
) -> str:
    automation_logger = configure_automation_logging(
        show_terminal_logs=show_terminal_logs
    )

    try:
        prepared_user = prepare_user_value(search_value)
        normalized_expiration_date = normalize_expiration_date(expiration_date)

        automation_logger.info(
            "Iniciando automação individual. Usuário: %s. "
            "Navegador visível: %s.",
            prepared_user,
            "sim" if mostrar_browser else "não"
        )

        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(
                headless=not mostrar_browser
            )
            context = browser.new_context()

            try:
                page = context.new_page()
                login_to_system(page, login, password)

                result = process_user_prepared_value(
                    page,
                    prepared_user,
                    normalized_expiration_date
                )

                automation_logger.info(
                    "Automação individual finalizada. Resultado: %s.",
                    result
                )

                return result
            finally:
                context.close()
                browser.close()
                automation_logger.info("Navegador encerrado.")

    except Exception:
        automation_logger.exception(
            "Automação individual finalizada com erro."
        )
        raise


# Executa a prorrogação unitária para uma lista de 1 a 5 usuários manuais.
def run_multi_automation(
    login: str,
    password: str,
    expiration_date: str,
    users: List[str],
    show_terminal_logs: bool = False,
    mostrar_browser: bool = True
) -> Tuple[str, str, str]:
    automation_logger = configure_automation_logging(
        show_terminal_logs=show_terminal_logs
    )

    messages: List[str] = []

    try:
        normalized_expiration_date = normalize_expiration_date(expiration_date)

        automation_logger.info(
            "Iniciando automação unitária multi-usuário. Total: %s. "
            "Navegador visível: %s.",
            len(users),
            "sim" if mostrar_browser else "não"
        )

        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(
                headless=not mostrar_browser
            )
            context = browser.new_context()

            try:
                page = context.new_page()
                login_to_system(page, login, password)

                for index, prepared_user in enumerate(users, start=1):
                    automation_logger.info(
                        "Processando usuário %s/%s: %s.",
                        index,
                        len(users),
                        prepared_user
                    )

                    try:
                        message = process_user_prepared_value(
                            page,
                            prepared_user,
                            normalized_expiration_date
                        )
                    except Exception as exc:
                        message = f"{ERROR_MESSAGE_PREFIX} {str(exc)}"

                        automation_logger.exception(
                            "Erro ao processar usuário %s.",
                            prepared_user
                        )

                    messages.append(message)

            finally:
                context.close()
                browser.close()
                automation_logger.info("Navegador encerrado.")

        if len(messages) == 1:
            mensagem = messages[0]
            detalhe = ""
        else:
            resumo = summarize_messages(messages)
            mensagem = f"Unitária concluída. {resumo}"
            detalhe = format_multi_results(users, messages)

        cor = result_color(messages)

        automation_logger.info(
            "Automação unitária multi-usuário finalizada. %s",
            mensagem
        )

        return mensagem, cor, detalhe

    except Exception:
        automation_logger.exception(
            "Automação unitária multi-usuário finalizada com erro."
        )
        raise


# Executa automação em lote.
def run_batch_automation(
    login: str,
    password: str,
    spreadsheet_path: str,
    report_directory: str,
    expiration_date: str,
    show_terminal_logs: bool = False,
    mostrar_browser: bool = True
) -> str:
    automation_logger = configure_automation_logging(
        show_terminal_logs=show_terminal_logs
    )

    try:
        normalized_expiration_date = normalize_expiration_date(expiration_date)

        automation_logger.info(
            "Iniciando automação em lote. Planilha: %s. "
            "Pasta relatório: %s. Navegador visível: %s.",
            spreadsheet_path,
            report_directory,
            "sim" if mostrar_browser else "não"
        )

        headers, source_rows = read_spreadsheet(spreadsheet_path)
        user_column = identify_user_column(headers)
        report_path = build_report_path(
            report_directory,
            spreadsheet_path
        )

        automation_logger.info(
            "Planilha carregada. Linhas úteis: %s. Coluna de usuários: %s.",
            len(source_rows),
            user_column
        )

        batch_rows: List[Row] = []

        for source_row in source_rows:
            user_value = source_row.get(user_column, "")
            report_user = normalize_user_value(user_value)

            batch_row: Row = {
                REPORT_COLUMN_USER: report_user,
                REPORT_COLUMN_NAME: "",
                BATCH_ENTRY_PREPARED_USER: "",
            }

            try:
                prepared_user = prepare_user_value(user_value)
                batch_row[REPORT_COLUMN_USER] = prepared_user
                batch_row[BATCH_ENTRY_PREPARED_USER] = prepared_user
            except ValueError as exc:
                batch_row[REPORT_COLUMN_NAME] = f"{ERROR_MESSAGE_PREFIX} {str(exc)}"

                automation_logger.warning(
                    "Usuário inválido no lote: %s. Erro: %s.",
                    report_user,
                    str(exc)
                )

            batch_rows.append(batch_row)

        prepared_rows = [
            row for row in batch_rows
            if row.get(BATCH_ENTRY_PREPARED_USER)
        ]

        automation_logger.info(
            "Usuários válidos para processamento: %s.",
            len(prepared_rows)
        )

        if prepared_rows:
            with sync_playwright() as playwright:
                browser = playwright.chromium.launch(
                    headless=not mostrar_browser
                )
                context = browser.new_context()

                try:
                    page = context.new_page()
                    login_to_system(page, login, password)

                    for index, batch_row in enumerate(
                        prepared_rows,
                        start=1
                    ):
                        prepared_user = batch_row.get(
                            BATCH_ENTRY_PREPARED_USER,
                            ""
                        )

                        automation_logger.info(
                            "Processando usuário %s/%s: %s.",
                            index,
                            len(prepared_rows),
                            prepared_user
                        )

                        try:
                            report_message = process_user_prepared_value(
                                page,
                                prepared_user,
                                normalized_expiration_date
                            )

                            automation_logger.info(
                                "Resultado do usuário %s: %s.",
                                prepared_user,
                                report_message
                            )
                        except Exception as exc:
                            report_message = f"{ERROR_MESSAGE_PREFIX} {str(exc)}"

                            automation_logger.exception(
                                "Erro ao processar usuário %s.",
                                prepared_user
                            )

                        batch_row[REPORT_COLUMN_NAME] = report_message

                finally:
                    context.close()
                    browser.close()
                    automation_logger.info("Navegador encerrado.")
        else:
            automation_logger.warning(
                "Nenhum usuário válido para processamento."
            )

        report_rows = [
            build_report_row(
                row.get(REPORT_COLUMN_USER, ""),
                row.get(REPORT_COLUMN_NAME, "")
            )
            for row in batch_rows
        ]

        write_report(
            spreadsheet_path,
            report_path,
            report_rows
        )

        resumo = summarize_batch_results(batch_rows)

        result = (
            f"Lote finalizado. {resumo} "
            f"Relatório: {report_path}"
        )

        automation_logger.info("Relatório gerado: %s.", report_path)
        automation_logger.info(result)

        return result

    except Exception:
        automation_logger.exception(
            "Automação em lote finalizada com erro."
        )
        raise