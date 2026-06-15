import ctypes
import os
import sys
import threading
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox

import customtkinter as ctk
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from playwright.sync_api import sync_playwright

sys.path.append(str(Path(__file__).resolve().parent.parent))

from PrinterAGHU import fazer_login, navegar_ate_modulo, processar_computadores
from autenticador import AGHU_URL

BASE_DIR = Path(__file__).resolve().parent
LOGS_DIR = BASE_DIR / "logs"


def esconder_console_windows() -> None:
    if os.name != "nt":
        return

    hwnd = ctypes.windll.kernel32.GetConsoleWindow()
    if hwnd:
        ctypes.windll.user32.ShowWindow(hwnd, 0)


def ler_planilha_entrada(caminho_planilha: str) -> pd.DataFrame:
    caminho = Path(caminho_planilha)

    if not caminho.exists():
        raise FileNotFoundError(f"Planilha de entrada não encontrada: {caminho}")

    extensao = caminho.suffix.lower()

    if extensao in {".xlsx", ".xlsm"}:
        df = pd.read_excel(caminho, dtype=str, engine="openpyxl")
    elif extensao == ".csv":
        try:
            df = pd.read_csv(caminho, sep=";", dtype=str, encoding="utf-8-sig")
        except UnicodeDecodeError:
            df = pd.read_csv(caminho, sep=";", dtype=str, encoding="latin1")
    else:
        raise ValueError("Formato inválido. Use .xlsx, .xlsm ou .csv.")

    df.columns = df.columns.str.strip()
    df = df.fillna("")

    colunas_obrigatorias = ["IPPC", "HostPrinter", "PrinterClass"]
    colunas_faltantes = [
        coluna for coluna in colunas_obrigatorias if coluna not in df.columns
    ]

    if colunas_faltantes:
        raise ValueError(
            "Planilha inválida. Colunas obrigatórias ausentes: "
            + ", ".join(colunas_faltantes)
        )

    return df


def normalizar_saida_xlsx(caminho_saida: str) -> Path:
    caminho = Path(caminho_saida).expanduser()

    if not caminho.suffix:
        caminho = caminho.with_suffix(".xlsx")

    if caminho.suffix.lower() != ".xlsx":
        raise ValueError("A planilha de saída deve ser um arquivo .xlsx.")

    caminho.parent.mkdir(parents=True, exist_ok=True)
    return caminho


def localizar_csv_relatorio_gerado(inicio_execucao: float) -> Path:
    arquivos = list(LOGS_DIR.glob("log_resultado_*.csv"))

    arquivos_recentes = [
        arquivo
        for arquivo in arquivos
        if arquivo.stat().st_mtime >= inicio_execucao - 2
    ]

    if arquivos_recentes:
        return max(arquivos_recentes, key=lambda arquivo: arquivo.stat().st_mtime)

    if arquivos:
        return max(arquivos, key=lambda arquivo: arquivo.stat().st_mtime)

    raise FileNotFoundError("Nenhum relatório CSV foi gerado na pasta logs.")


def converter_csv_relatorio_para_xlsx(caminho_csv: Path, caminho_xlsx: Path) -> None:
    with caminho_csv.open("r", encoding="utf-8-sig") as arquivo:
        linha_auditoria = arquivo.readline().strip()

    df_relatorio = pd.read_csv(
        caminho_csv,
        sep=";",
        dtype=str,
        encoding="utf-8-sig",
        skiprows=1,
    ).fillna("")

    with pd.ExcelWriter(caminho_xlsx, engine="openpyxl") as writer:
        df_relatorio.to_excel(
            writer,
            index=False,
            sheet_name="Relatório",
            startrow=3,
        )

        worksheet = writer.sheets["Relatório"]

        worksheet["A1"] = linha_auditoria or "Atualizado por:"
        worksheet["A2"] = f"Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
        worksheet["A1"].font = Font(bold=True)
        worksheet["A2"].font = Font(italic=True)

        header_row = 4
        first_data_row = 5
        max_row = max(worksheet.max_row, first_data_row)
        max_col = max(worksheet.max_column, 1)
        last_col_letter = get_column_letter(max_col)

        worksheet.freeze_panes = "A5"
        worksheet.auto_filter.ref = f"A{header_row}:{last_col_letter}{max_row}"

        header_fill = PatternFill("solid", fgColor="D9EAF7")

        for cell in worksheet[header_row]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for column_index in range(1, max_col + 1):
            column_letter = get_column_letter(column_index)
            max_length = 0

            for cell in worksheet[column_letter]:
                valor = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(valor))

            worksheet.column_dimensions[column_letter].width = min(max_length + 2, 60)


def executar_automacao_aghu(
    usuario: str,
    senha: str,
    mostrar_console: bool,
    mostrar_browser: bool,
    caminho_planilha_entrada: str,
    caminho_planilha_saida: str,
) -> str:
    if not usuario or not senha:
        raise ValueError("Preencha usuário de rede e senha.")

    if not mostrar_console:
        esconder_console_windows()

    os.chdir(BASE_DIR)

    planilha = ler_planilha_entrada(caminho_planilha_entrada)
    caminho_saida = normalizar_saida_xlsx(caminho_planilha_saida)
    LOGS_DIR.mkdir(parents=True, exist_ok=True)

    inicio_execucao = datetime.now().timestamp()

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(
            headless=not mostrar_browser,
            slow_mo=500,
        )
        context = browser.new_context(ignore_https_errors=True)
        page = context.new_page()

        try:
            page.goto(AGHU_URL)
            print(f"🌍 Ambiente acessado: {page.url}")

            fazer_login(page, usuario, senha)
            page, janela_sistema = navegar_ate_modulo(context, page, usuario, senha)

            processar_computadores(
                context,
                page,
                janela_sistema,
                planilha,
                usuario,
                senha,
            )
        finally:
            browser.close()

    caminho_csv = localizar_csv_relatorio_gerado(inicio_execucao)
    converter_csv_relatorio_para_xlsx(caminho_csv, caminho_saida)

    return f"Processo concluído. Relatório salvo em: {caminho_saida}"


class AghuPrinterApp(ctk.CTk):
    def __init__(self):
        super().__init__()

        self.title("AGHUX Bot - Impressoras")
        self.geometry("760x520")
        self.resizable(False, False)
        self.grid_columnconfigure(0, weight=1)

        self.var_browser = tk.BooleanVar(value=True)
        self.var_console = tk.BooleanVar(value=True)

        self.label_title = ctk.CTkLabel(
            self,
            text="AGHUX Bot - Impressoras",
            font=ctk.CTkFont(size=20, weight="bold"),
        )
        self.label_title.grid(row=0, column=0, padx=20, pady=(28, 16))

        self.frame_inputs = ctk.CTkFrame(self)
        self.frame_inputs.grid(row=1, column=0, padx=20, pady=10, sticky="ew")
        self.frame_inputs.grid_columnconfigure(1, weight=1)

        self.create_login_fields()
        self.create_execution_options()
        self.create_spreadsheet_fields()

        self.button_run = ctk.CTkButton(
            self,
            text="Iniciar Robô",
            command=self.start_automation,
            font=ctk.CTkFont(weight="bold"),
            height=38,
        )
        self.button_run.grid(row=2, column=0, padx=20, pady=(24, 12))

        self.label_status = ctk.CTkLabel(
            self,
            text="Pronto para execução.",
            text_color="gray",
            wraplength=680,
        )
        self.label_status.grid(row=3, column=0, padx=20, pady=(8, 16))

    def create_login_fields(self) -> None:
        self.label_user = ctk.CTkLabel(self.frame_inputs, text="Usuário de Rede:")
        self.label_user.grid(row=0, column=0, padx=12, pady=(16, 8), sticky="e")

        self.entry_user = ctk.CTkEntry(
            self.frame_inputs,
            placeholder_text="Digite seu usuário de rede",
        )
        self.entry_user.grid(
            row=0,
            column=1,
            columnspan=2,
            padx=12,
            pady=(16, 8),
            sticky="ew",
        )
        self.entry_user.focus()

        self.label_password = ctk.CTkLabel(self.frame_inputs, text="Senha:")
        self.label_password.grid(row=1, column=0, padx=12, pady=8, sticky="e")

        self.entry_password = ctk.CTkEntry(
            self.frame_inputs,
            placeholder_text="Digite sua senha",
            show="*",
        )
        self.entry_password.grid(
            row=1,
            column=1,
            columnspan=2,
            padx=12,
            pady=8,
            sticky="ew",
        )

    def create_execution_options(self) -> None:
        self.checkbox_browser = ctk.CTkCheckBox(
            self.frame_inputs,
            text="Exibir Navegador (Modo Visual)",
            variable=self.var_browser,
            command=self.validate_visibility_options_from_browser,
        )
        self.checkbox_browser.grid(
            row=2,
            column=1,
            columnspan=2,
            padx=12,
            pady=(14, 6),
            sticky="w",
        )

        self.checkbox_console = ctk.CTkCheckBox(
            self.frame_inputs,
            text="Exibir Terminal de processos (logs)",
            variable=self.var_console,
            command=self.validate_visibility_options_from_console,
        )
        self.checkbox_console.grid(
            row=3,
            column=1,
            columnspan=2,
            padx=12,
            pady=(6, 14),
            sticky="w",
        )

    def create_spreadsheet_fields(self) -> None:
        self.label_spreadsheet_in = ctk.CTkLabel(
            self.frame_inputs,
            text="Planilha entrada:",
        )
        self.label_spreadsheet_in.grid(row=4, column=0, padx=12, pady=10, sticky="e")

        self.entry_spreadsheet_in = ctk.CTkEntry(
            self.frame_inputs,
            placeholder_text="Arquivo .xlsx, .xlsm ou .csv",
        )
        self.entry_spreadsheet_in.grid(
            row=4,
            column=1,
            padx=12,
            pady=10,
            sticky="ew",
        )

        self.button_select_spreadsheet_in = ctk.CTkButton(
            self.frame_inputs,
            text="Selecionar",
            width=110,
            command=self.select_input_spreadsheet,
        )
        self.button_select_spreadsheet_in.grid(row=4, column=2, padx=12, pady=10)

        self.label_spreadsheet_out = ctk.CTkLabel(
            self.frame_inputs,
            text="Planilha saída:",
        )
        self.label_spreadsheet_out.grid(
            row=5,
            column=0,
            padx=12,
            pady=(10, 16),
            sticky="e",
        )

        self.entry_spreadsheet_out = ctk.CTkEntry(
            self.frame_inputs,
            placeholder_text="Arquivo .xlsx do relatório",
        )
        self.entry_spreadsheet_out.grid(
            row=5,
            column=1,
            padx=12,
            pady=(10, 16),
            sticky="ew",
        )

        self.button_select_spreadsheet_out = ctk.CTkButton(
            self.frame_inputs,
            text="Selecionar",
            width=110,
            command=self.select_output_spreadsheet,
        )
        self.button_select_spreadsheet_out.grid(
            row=5,
            column=2,
            padx=12,
            pady=(10, 16),
        )

    def validate_visibility_options_from_browser(self) -> None:
        if not self.var_browser.get() and not self.var_console.get():
            self.var_browser.set(True)
            messagebox.showwarning(
                "Ação bloqueada",
                "Para evitar processos invisíveis, mantenha o navegador ou o terminal ativo.",
            )

    def validate_visibility_options_from_console(self) -> None:
        if not self.var_console.get() and not self.var_browser.get():
            self.var_console.set(True)
            messagebox.showwarning(
                "Ação bloqueada",
                "Para evitar processos invisíveis, mantenha o navegador ou o terminal ativo.",
            )

    def select_input_spreadsheet(self) -> None:
        file_path = filedialog.askopenfilename(
            title="Selecione a planilha de entrada",
            filetypes=(
                ("Planilhas", "*.xlsx *.xlsm *.csv"),
                ("Excel", "*.xlsx *.xlsm"),
                ("CSV", "*.csv"),
            ),
        )

        if file_path:
            self.entry_spreadsheet_in.delete(0, "end")
            self.entry_spreadsheet_in.insert(0, file_path)

            if not self.entry_spreadsheet_out.get().strip():
                default_output = Path(file_path).with_name(
                    f"relatorio_aghu_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
                )
                self.entry_spreadsheet_out.insert(0, str(default_output))

    def select_output_spreadsheet(self) -> None:
        default_name = f"relatorio_aghu_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

        file_path = filedialog.asksaveasfilename(
            title="Salvar relatório como",
            defaultextension=".xlsx",
            initialfile=default_name,
            filetypes=(("Excel", "*.xlsx"),),
        )

        if file_path:
            self.entry_spreadsheet_out.delete(0, "end")
            self.entry_spreadsheet_out.insert(0, file_path)

    def start_automation(self) -> None:
        usuario = self.entry_user.get().strip()
        senha = self.entry_password.get().strip()
        caminho_entrada = self.entry_spreadsheet_in.get().strip()
        caminho_saida = self.entry_spreadsheet_out.get().strip()
        mostrar_console = bool(self.var_console.get())
        mostrar_browser = bool(self.var_browser.get())

        if not usuario or not senha:
            self.show_status("Erro: preencha usuário de rede e senha.", "red")
            return

        if not caminho_entrada:
            self.show_status("Erro: informe a planilha de entrada.", "red")
            return

        if not caminho_saida:
            self.show_status("Erro: informe a planilha de saída do relatório.", "red")
            return

        try:
            normalizar_saida_xlsx(caminho_saida)
        except Exception as exc:
            self.show_status(f"Erro: {exc}", "red")
            return

        self.show_status("Iniciando automação...", "blue")
        self.button_run.configure(state="disabled", text="Executando...")

        thread = threading.Thread(
            target=self.run_playwright_task,
            args=(
                usuario,
                senha,
                mostrar_console,
                mostrar_browser,
                caminho_entrada,
                caminho_saida,
            ),
            daemon=True,
        )
        thread.start()

    def run_playwright_task(
        self,
        usuario: str,
        senha: str,
        mostrar_console: bool,
        mostrar_browser: bool,
        caminho_entrada: str,
        caminho_saida: str,
    ) -> None:
        try:
            result_msg = executar_automacao_aghu(
                usuario=usuario,
                senha=senha,
                mostrar_console=mostrar_console,
                mostrar_browser=mostrar_browser,
                caminho_planilha_entrada=caminho_entrada,
                caminho_planilha_saida=caminho_saida,
            )
            self.after(0, self.finish_automation, result_msg, "green")
        except Exception as exc:
            self.after(0, self.finish_automation, f"Erro: {exc}", "red")

    def finish_automation(self, message: str, color: str) -> None:
        self.show_status(message, color)
        self.button_run.configure(state="normal", text="Iniciar Robô")

    def show_status(self, message: str, color: str) -> None:
        self.label_status.configure(text=message, text_color=color)


if __name__ == "__main__":
    ctk.set_appearance_mode("System")
    ctk.set_default_color_theme("blue")

    app = AghuPrinterApp()
    app.mainloop()