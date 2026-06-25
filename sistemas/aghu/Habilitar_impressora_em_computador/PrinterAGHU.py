import os
import re
import sys
import time
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Mapping, Optional, Sequence, Tuple

sys.path.append(str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from playwright.sync_api import BrowserContext, Page

#Imports de classes utilitárias públicas.
from autenticador import AGHU_URL, autenticar_aghu_page, exigir_login_valido
from menu import navegar_menu_aghu



from AddPrinterAGHU import (
    cadastrar_nova_impressora,
    consultar_dados_site_secundario,
    navegar_ate_cadastro_impressora,
)


BASE_DIR = Path(__file__).resolve().parent
SUPPORTED_EXTENSIONS = (".xlsx",)
Row = Dict[str, Any]

# ==========================================
# CAMINHOS DO PROCEDIMENTO
# ==========================================

CAMINHO_MENU_IMPRESSORA_POR_COMPUTADOR = (
    "Outros Módulos",
    "Configuração",
    "Impressão",
    "Cadastros",
    "Impressora por Computador",
)
COLUNAS_OBRIGATORIAS_PLANILHA = ["IPPC", "HostPrinter", "PrinterClass"]
COLUNAS_RELATORIO = (
    "HostPC",
    "IPPC",
    "HostPrinter",
    "IPPrinter",
    "PrinterClass",
    "Status",
    "Detalhes",
)
ALIASES_COLUNAS_PLANILHA = {
    "HostPC": (
        "HostPC",
        "Host PC",
        "Computador",
        "Nome Computador",
        "Nome do Computador",
    ),
    "IPPC": (
        "IPPC",
        "IP PC",
        "IP do PC",
        "IP Computador",
        "IP do Computador",
        "Endereco IP PC",
        "Endereco IP do PC",
    ),
    "HostPrinter": (
        "HostPrinter",
        "Host Printer",
        "Impressora",
        "Fila",
        "Fila Impressora",
        "Fila da Impressora",
        "Nome Impressora",
        "Nome da Impressora",
    ),
    "IPPrinter": (
        "IPPrinter",
        "IP Printer",
        "IP da Impressora",
        "IP Impressora",
        "Endereco IP Impressora",
        "Endereco IP da Impressora",
    ),
    "PrinterClass": (
        "PrinterClass",
        "Printer Class",
        "Classe",
        "Classe Impressao",
        "Classe de Impressao",
        "Tipo",
        "Tipo CUPS",
    ),
}
CARACTERES_DE_VALOR = r"A-Za-z0-9_.-"
TABELA_COMPUTADOR_IMPRESSORA_SELECTOR = (
    '[id="tabelaComputadorImpressora:resultList_data"]'
)
MENSAGEM_ERRO_PESQUISA_INDEFINIDA = (
    "Conferir manualmente: pesquisa nao retornou nem linhas nem mensagem de "
    "nenhum registro encontrado."
)
MAX_TENTATIVAS_AUTENTICAR_NOVA_ABA = 3
INTERVALO_RETRY_AUTENTICAR_NOVA_ABA_SEGUNDOS = 1
MAX_TENTATIVAS_PROCESSAMENTO_LINHA = 3
MAX_TENTATIVAS_ESTOQUE = 3

# Layout XLSX
XLSX_HEADER_ROW = 1
XLSX_FREEZE_PANES_CELL = "A2"
XLSX_MIN_COLUMN_WIDTH = 12
XLSX_MAX_COLUMN_WIDTH = 60
XLSX_COLUMN_PADDING = 2


@dataclass(frozen=True)
class DadosLinhaPlanilha:
    ip_pc: str
    impressora_alvo: str
    classe_impressao: str


@dataclass(frozen=True)
class ResultadoLinha:
    status: str
    detalhes: str


@dataclass(frozen=True)
class AcaoVinculo:
    tipo: str
    registro: dict | None = None


class FalhaTecnicaProcessamento(RuntimeError):
    def __init__(self, passo: str, erro: Exception):
        super().__init__(str(erro))
        self.passo = passo
        self.erro = erro


def _normalizar_busca(valor: object) -> str:
    return re.sub(r"\s+", " ", str(valor or "").strip()).casefold()


def _normalizar_texto_simples(valor: object) -> str:
    return _normalizar_busca(valor)


def _criar_regex_valor_exato(valor: object, flags: int = re.IGNORECASE) -> re.Pattern:
    valor_normalizado = _normalizar_busca(valor)
    return re.compile(
        rf"(?<![{CARACTERES_DE_VALOR}])"
        rf"{re.escape(valor_normalizado)}"
        rf"(?![{CARACTERES_DE_VALOR}])",
        flags,
    )


def _contem_valor_exato(texto: object, valor: object) -> bool:
    valor_normalizado = _normalizar_busca(valor)

    if not valor_normalizado:
        return False

    texto_normalizado = _normalizar_busca(texto)
    padrao = _criar_regex_valor_exato(valor_normalizado, flags=0)

    return bool(padrao.search(texto_normalizado))


def _valor_exato(valor_atual: object, valor_esperado: object) -> bool:
    return _normalizar_busca(valor_atual) == _normalizar_busca(valor_esperado)


def _regex_ip_celula(ip_pc: str) -> re.Pattern:
    return re.compile(rf"^\s*{re.escape(str(ip_pc).strip())}\s*$")


def _ip_celula_confere(valor_celula: object, ip_pc: str) -> bool:
    return bool(_regex_ip_celula(ip_pc).match(str(valor_celula or "")))


def _classe_impressao_aghu(tipo_cups: str) -> str:
    if _valor_exato(tipo_cups, "PDF"):
        return "A"

    return str(tipo_cups or "").strip()


def _tbody_resultados(janela_sistema):
    return janela_sistema.locator(TABELA_COMPUTADOR_IMPRESSORA_SELECTOR)


def _linhas_resultado(tbody):
    return tbody.locator("> tr[data-ri]")


def _linha_vazia_resultado(tbody):
    return tbody.locator(
        "> tr.ui-datatable-empty-message",
        has_text="Nenhum registro encontrado!",
    )


def _texto_celula(linha_tabela, indice: int) -> str:
    try:
        celulas = linha_tabela.locator("td")
        if indice >= celulas.count():
            return ""

        return celulas.nth(indice).inner_text(timeout=1000).strip()
    except Exception:
        return ""


def _aguardar_estado_resultado_pesquisa(
    janela_sistema,
    timeout_ms: int = 7000,
) -> tuple[str, object]:
    tbody = _tbody_resultados(janela_sistema)
    linhas = _linhas_resultado(tbody)
    linha_vazia = _linha_vazia_resultado(tbody)
    fim = time.monotonic() + (timeout_ms / 1000)

    while time.monotonic() < fim:
        try:
            if linhas.count() > 0 and linhas.first.is_visible(timeout=250):
                return "linhas", linhas
        except Exception:
            pass

        try:
            if linha_vazia.is_visible(timeout=250):
                return "vazio", linhas
        except Exception:
            pass

        time.sleep(0.15)

    return "indefinido", linhas


def _registro_confere_tipo_e_classe(registro: dict, tipo_cups_esperado: str) -> bool:
    if not _valor_exato(registro.get("tipo_cups", ""), tipo_cups_esperado):
        return False

    classe_esperada = _classe_impressao_aghu(tipo_cups_esperado)

    if not classe_esperada:
        return True

    return _valor_exato(registro.get("classe", ""), classe_esperada)


def _registro_eh_pdf(registro: dict) -> bool:
    return _registro_confere_tipo_e_classe(registro, "PDF")


def _decidir_acao_linhas(
    registros_linhas: list[dict],
    impressora_alvo: str,
    classe_impressao: str,
) -> tuple[str, dict | None]:
    linha_pdf = None

    for registro in registros_linhas:
        if _valor_exato(registro.get("fila", ""), impressora_alvo):
            if _registro_confere_tipo_e_classe(registro, classe_impressao):
                return "mantido", registro

            return "conferir", registro

        if _registro_eh_pdf(registro):
            linha_pdf = registro

    if linha_pdf is not None:
        return "alterar", linha_pdf

    return "incluir", None


def _coletar_linhas_computador(janela_sistema, ip_pc: str) -> tuple[str, list[dict]]:
    estado_pesquisa, linhas_tabela = _aguardar_estado_resultado_pesquisa(janela_sistema)
    registros_linhas = []

    if estado_pesquisa != "linhas":
        return estado_pesquisa, registros_linhas

    for indice in range(linhas_tabela.count()):
        linha_tabela = linhas_tabela.nth(indice)

        try:
            if not linha_tabela.is_visible(timeout=1000):
                continue

            texto_linha = linha_tabela.inner_text(timeout=1000)
        except Exception:
            continue

        ip_linha = _texto_celula(linha_tabela, 1)

        if not _ip_celula_confere(ip_linha, ip_pc):
            continue

        registros_linhas.append(
            {
                "linha": linha_tabela,
                "texto": texto_linha,
                "ip": ip_linha,
                "computador": _texto_celula(linha_tabela, 2),
                "descricao": _texto_celula(linha_tabela, 3),
                "classe": _texto_celula(linha_tabela, 4),
                "fila": _texto_celula(linha_tabela, 5),
                "tipo_cups": _texto_celula(linha_tabela, 6),
            }
        )

    return estado_pesquisa, registros_linhas


def _extrair_ips(texto: object) -> list[str]:
    return re.findall(
        r"(?<!\S)(\d{1,3}(?:\.\d{1,3}){3})(?!\S)",
        str(texto or ""),
    )


def _validar_computador_selecionado(
    campo_computador,
    ip_pc: str,
    texto_item_selecionado: str,
) -> tuple[bool, str]:
    try:
        valor_campo = campo_computador.input_value(timeout=1000).strip()
    except Exception:
        valor_campo = ""

    textos = [texto for texto in (valor_campo, texto_item_selecionado) if texto]
    ips_encontrados = []

    for texto in textos:
        ips_encontrados.extend(_extrair_ips(texto))

    if ips_encontrados:
        ips_divergentes = [ip for ip in ips_encontrados if ip != ip_pc]
        if not ips_divergentes and ip_pc in ips_encontrados:
            return True, ""

        return False, ips_divergentes[0] if ips_divergentes else ips_encontrados[0]

    if _ip_celula_confere(valor_campo, ip_pc):
        return True, ""

    return False, valor_campo or texto_item_selecionado or "nao identificado"


def _mensagem_dialog(janela_sistema, seletor: str):
    return janela_sistema.locator(
        f'#msgDialog[aria-hidden="false"] #messagesInDialog {seletor}'
    )


def _aguardar_resultado_gravacao(
    janela_sistema,
    page: Page | None = None,
    timeout_ms: int = 10000,
) -> tuple[str, str]:
    containers = [janela_sistema]

    if page is not None:
        containers.append(page)

    mensagens_sucesso = [
        _mensagem_dialog(container, "span.ui-messages-info-summary")
        for container in containers
    ]
    mensagens_erro = [
        _mensagem_dialog(container, "span.ui-messages-error-summary")
        for container in containers
    ]
    fim = time.monotonic() + (timeout_ms / 1000)

    while time.monotonic() < fim:
        for mensagem_erro in mensagens_erro:
            try:
                if mensagem_erro.count() > 0 and mensagem_erro.first.is_visible(timeout=250):
                    return "erro", mensagem_erro.first.inner_text(timeout=1000).strip()
            except Exception:
                pass

        for mensagem_sucesso in mensagens_sucesso:
            try:
                if mensagem_sucesso.count() > 0 and mensagem_sucesso.first.is_visible(timeout=250):
                    return "sucesso", mensagem_sucesso.first.inner_text(timeout=1000).strip()
            except Exception:
                pass

        time.sleep(0.15)

    return "indefinido", ""


def _erro_classe_pdf_duplicada(mensagem: str) -> bool:
    mensagem_normalizada = _normalizar_texto_simples(mensagem)
    return (
        "existe uma impressora cadastrada" in mensagem_normalizada
        and "classe a" in mensagem_normalizada
    )


def _aguardar_botao_pesquisar_se_possivel(janela_sistema) -> None:
    try:
        janela_sistema.get_by_role("button", name="Pesquisar").wait_for(
            state="visible",
            timeout=3000,
        )
    except Exception:
        pass


def _limpar_estado_formulario(janela_sistema, page: Page | None = None) -> None:
    containers = [janela_sistema]

    if page is not None:
        containers.append(page)

    for container in containers:
        try:
            container.locator(
                '#msgDialog[aria-hidden="false"] a.ui-dialog-titlebar-close, '
                '#msgDialog[aria-hidden="false"] .ui-dialog-titlebar-close'
            ).first.click(timeout=1000)
            break
        except Exception:
            pass

    try:
        janela_sistema.get_by_role("button", name="Cancelar").click(timeout=1000)
    except Exception:
        pass

    try:
        janela_sistema.locator("button:has(.aghu-icon-cleaner-aghu)").first.click(timeout=1500)
    except Exception:
        pass

def _valor_planilha_em_branco(valor: object) -> bool:
    if valor is None:
        return True

    try:
        if valor != valor:
            return True
    except Exception:
        pass

    return str(valor).strip() == ""


def _campos_obrigatorios_planilha_em_branco(
    linha: Mapping[str, Any],
) -> list[str]:
    return [
        coluna
        for coluna in COLUNAS_OBRIGATORIAS_PLANILHA
        if _valor_planilha_em_branco(linha.get(coluna, ""))
    ]


def normalize_column_name(value: Any) -> str:
    text = "" if value is None else str(value).strip()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(
        char for char in text
        if not unicodedata.combining(char)
    )
    return re.sub(r"\s+", " ", text).casefold()


def validate_spreadsheet_extension(spreadsheet_path: str) -> str:
    extension = Path(spreadsheet_path).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(SUPPORTED_EXTENSIONS)
        raise ValueError(
            f"Formato de planilha nao suportado: {extension}. "
            f"Use apenas: {supported}."
        )

    return extension


def generate_report_filename(
    extension: str,
    now: Optional[datetime] = None,
) -> str:
    reference_date = now or datetime.now()
    clean_extension = extension.lower().strip()

    if not clean_extension.startswith("."):
        clean_extension = f".{clean_extension}"

    timestamp = reference_date.strftime("%d_%m_%y_%Hh%Mm%S")
    return f"Resultado_{timestamp}{clean_extension}"


def get_available_report_path(report_path: Path) -> Path:
    if not report_path.exists():
        return report_path

    counter = 2
    parent = report_path.parent
    stem = report_path.stem
    suffix = report_path.suffix

    while True:
        candidate = parent / f"{stem}_{counter}{suffix}"

        if not candidate.exists():
            return candidate

        counter += 1


def build_report_path(
    report_directory: str,
    source_spreadsheet_path: str,
    now: Optional[datetime] = None,
) -> Path:
    report_dir = Path(report_directory).expanduser()

    if not report_dir.exists():
        raise FileNotFoundError(
            f"A pasta de relatorio nao existe: {report_dir}"
        )

    if not report_dir.is_dir():
        raise NotADirectoryError(
            f"O caminho de relatorio nao e uma pasta: {report_dir}"
        )

    extension = validate_spreadsheet_extension(source_spreadsheet_path)
    report_filename = generate_report_filename(extension, now)
    return get_available_report_path(report_dir / report_filename)


def read_spreadsheet(spreadsheet_path: str) -> Tuple[list[str], list[Row]]:
    source_path = Path(spreadsheet_path).expanduser()

    if not source_path.exists():
        raise FileNotFoundError(f"Planilha nao encontrada: {source_path}")

    validate_spreadsheet_extension(str(source_path))
    return read_xlsx(source_path)


def read_xlsx(source_path: Path) -> Tuple[list[str], list[Row]]:
    workbook = load_workbook(source_path, data_only=True)
    worksheet = workbook.active

    if worksheet is None:
        raise ValueError("A planilha nao possui aba ativa.")

    first_row = next(
        worksheet.iter_rows(min_row=1, max_row=1),
        None,
    )

    if first_row is None:
        raise ValueError("A planilha XLSX esta vazia.")

    raw_headers = [cell.value for cell in first_row]
    headers = make_unique_headers(raw_headers)
    rows: list[Row] = []

    for raw_row in worksheet.iter_rows(min_row=2, values_only=True):
        if is_empty_row(raw_row):
            continue

        rows.append(build_row(headers, raw_row))

    return headers, rows


def make_unique_headers(raw_headers: Sequence[Any]) -> list[str]:
    if not raw_headers:
        raise ValueError("A planilha nao possui cabecalho.")

    headers: list[str] = []
    seen: dict[str, int] = {}

    for index, raw_header in enumerate(raw_headers, start=1):
        header = "" if raw_header is None else str(raw_header).strip()

        if not header:
            header = f"coluna_{index}"

        if header in seen:
            seen[header] += 1
            header = f"{header}_{seen[header]}"
        else:
            seen[header] = 1

        headers.append(header)

    return headers


def build_row(headers: Sequence[str], raw_row: Sequence[Any]) -> Row:
    row: Row = {}

    for index, header in enumerate(headers):
        row[header] = raw_row[index] if index < len(raw_row) else ""

    return row


def is_empty_row(raw_row: Sequence[Any]) -> bool:
    return all(_valor_planilha_em_branco(value) for value in raw_row)


def identify_column_by_alias(
    headers: Sequence[str],
    aliases: Sequence[str],
    friendly_name: str,
) -> str:
    candidates = {
        normalize_column_name(alias)
        for alias in aliases
    }

    for header in headers:
        if normalize_column_name(header) in candidates:
            return header

    expected_columns = ", ".join(aliases)
    raise ValueError(
        f"Coluna obrigatoria nao encontrada para {friendly_name}. "
        f"A planilha deve conter uma destas colunas: {expected_columns}."
    )


def identify_spreadsheet_columns(headers: Sequence[str]) -> dict[str, str]:
    column_map: dict[str, str] = {}

    for canonical_name in COLUNAS_OBRIGATORIAS_PLANILHA:
        column_map[canonical_name] = identify_column_by_alias(
            headers,
            ALIASES_COLUNAS_PLANILHA[canonical_name],
            canonical_name,
        )

    for canonical_name in ("HostPC", "IPPrinter"):
        try:
            column_map[canonical_name] = identify_column_by_alias(
                headers,
                ALIASES_COLUNAS_PLANILHA[canonical_name],
                canonical_name,
            )
        except ValueError:
            continue

    return column_map


def normalize_spreadsheet_cell(value: Any) -> Any:
    if _valor_planilha_em_branco(value):
        return ""

    return value


def build_spreadsheet_row(
    source_row: Mapping[str, Any],
    column_map: Mapping[str, str],
) -> Row:
    row: Row = {}

    for canonical_name in (
        "HostPC",
        "IPPC",
        "HostPrinter",
        "IPPrinter",
        "PrinterClass",
    ):
        source_header = column_map.get(canonical_name)
        row[canonical_name] = (
            normalize_spreadsheet_cell(source_row.get(source_header, ""))
            if source_header is not None
            else ""
        )

    return row


def ler_planilha(caminho_arquivo: str) -> list[Row]:
    headers, source_rows = read_spreadsheet(caminho_arquivo)
    column_map = identify_spreadsheet_columns(headers)

    return [
        build_spreadsheet_row(source_row, column_map)
        for source_row in source_rows
    ]


def get_report_headers() -> list[str]:
    return list(COLUNAS_RELATORIO)


def write_report(
    source_spreadsheet_path: str,
    report_path: Path,
    rows: Sequence[Mapping[str, Any]],
) -> None:
    validate_spreadsheet_extension(source_spreadsheet_path)
    write_xlsx_report(report_path, get_report_headers(), rows)


def write_xlsx_report(
    report_path: Path,
    headers: Sequence[str],
    rows: Sequence[Mapping[str, Any]],
) -> None:
    workbook = Workbook()
    worksheet = workbook.active

    if worksheet is None:
        raise ValueError("A planilha nao possui aba ativa.")

    worksheet.title = "Relatorio"
    worksheet.append(list(headers))

    for row in rows:
        worksheet.append([
            row.get(header, "")
            for header in headers
        ])

    apply_xlsx_report_layout(worksheet)
    workbook.save(report_path)


def apply_xlsx_report_layout(worksheet: Worksheet) -> None:
    if worksheet.max_row < XLSX_HEADER_ROW or worksheet.max_column < 1:
        return

    worksheet.freeze_panes = XLSX_FREEZE_PANES_CELL
    worksheet.auto_filter.ref = build_xlsx_filter_range(worksheet)
    autofit_xlsx_columns(worksheet)


def build_xlsx_filter_range(worksheet: Worksheet) -> str:
    last_column = get_column_letter(worksheet.max_column)
    return f"A{XLSX_HEADER_ROW}:{last_column}{worksheet.max_row}"


def autofit_xlsx_columns(worksheet: Worksheet) -> None:
    for column_cells in worksheet.columns:
        if not column_cells:
            continue

        col_idx = column_cells[0].column

        if not isinstance(col_idx, int):
            continue

        column_letter = get_column_letter(col_idx)
        max_length = max(
            get_xlsx_cell_text_length(cell.value)
            for cell in column_cells
        )
        width = max_length + XLSX_COLUMN_PADDING
        worksheet.column_dimensions[column_letter].width = min(
            max(width, XLSX_MIN_COLUMN_WIDTH),
            XLSX_MAX_COLUMN_WIDTH,
        )


def get_xlsx_cell_text_length(value: Any) -> int:
    if value is None:
        return 0

    lines = str(value).splitlines() or [""]
    return max(len(line) for line in lines)


def fazer_login(
    page: Page,
    usuario_str: str,
    senha_str: str,
    *,
    url_aghu: str = AGHU_URL,
):
    print(f"Checando autenticação no AGHUX com o usuário: {usuario_str}")

    resultado = autenticar_aghu_page(
        page=page,
        usuario=usuario_str,
        senha=senha_str,
        url_login=url_aghu,
        timeout_ms=15000,
    )

    if resultado.status == "sessao_ativa":
        print("Sessão já estava ativa.")
    elif resultado.status == "sucesso":
        print("Login efetuado com sucesso.")
    else:
        print(f"Falha de autenticação: {resultado.mensagem}")

    exigir_login_valido(resultado)
    return resultado


def _fechar_page_silenciosamente(page: Page | None) -> None:
    if page is None:
        return

    try:
        if not page.is_closed():
            page.close()
    except Exception:
        pass


def _fechar_abas_contexto(
    context: BrowserContext,
    *,
    exceto: Page | None = None,
) -> None:
    for page_contexto in list(context.pages):
        if exceto is not None and page_contexto == exceto:
            continue

        _fechar_page_silenciosamente(page_contexto)


# ==========================================
# ISOLAMENTO DE SESSÃO (CLEAN STATE)
# ==========================================
def trocar_aba_aghux(
    context: BrowserContext,
    page_atual: Page,
    usuario_str: str,
    senha_str: str,
    *,
    url_aghu: str = AGHU_URL,
    max_tentativas_autenticacao: int = MAX_TENTATIVAS_AUTENTICAR_NOVA_ABA,
) -> Page:
    print("[Clean State] Fechando abas atuais e abrindo nova aba limpa.")

    _fechar_page_silenciosamente(page_atual)
    _fechar_abas_contexto(context)

    total_tentativas = max(1, max_tentativas_autenticacao)
    ultima_falha = "motivo nao identificado"

    for tentativa in range(1, total_tentativas + 1):
        nova_page: Page | None = None

        try:
            if tentativa > 1:
                print(
                    "[Clean State] Nova tentativa de autenticacao em aba limpa "
                    f"({tentativa}/{total_tentativas})."
                )

            nova_page = context.new_page()
            nova_page.goto(
                url_aghu,
                wait_until="domcontentloaded",
                timeout=15000,
            )
            print(f"Ambiente acessado: {nova_page.url}")

            resultado = autenticar_aghu_page(
                page=nova_page,
                usuario=usuario_str,
                senha=senha_str,
                url_login=url_aghu,
                timeout_ms=15000,
            )

            if resultado.status == "sessao_ativa":
                print("Sessao reaproveitada na nova aba.")
                _fechar_abas_contexto(context, exceto=nova_page)
                return nova_page

            if resultado.status == "sucesso":
                print("Login efetuado na nova aba.")
                _fechar_abas_contexto(context, exceto=nova_page)
                return nova_page

            ultima_falha = resultado.mensagem
            print(
                "Falha ao autenticar nova aba "
                f"({tentativa}/{total_tentativas}): {resultado.mensagem}"
            )

        except Exception as exc:
            ultima_falha = str(exc)
            print(
                "Falha ao preparar/autenticar nova aba "
                f"({tentativa}/{total_tentativas}): {ultima_falha}"
            )

        _fechar_page_silenciosamente(nova_page)

        if tentativa < total_tentativas:
            time.sleep(INTERVALO_RETRY_AUTENTICAR_NOVA_ABA_SEGUNDOS)

    _fechar_abas_contexto(context)
    raise RuntimeError(
        "Falha ao autenticar nova aba apos "
        f"{total_tentativas} tentativas: {ultima_falha}"
    )


def navegar_ate_modulo(
    context: BrowserContext,
    page_atual: Page,
    usuario_str: str,
    senha_str: str,
    *,
    url_aghu: str = AGHU_URL,
):
    print("🗺️ Navegando até o módulo de Impressora por Computador...")
    page = page_atual

    for tentativa in range(2):
        try:
            janela_sistema = navegar_menu_aghu(
                page=page,
                caminho=CAMINHO_MENU_IMPRESSORA_POR_COMPUTADOR,
            )
            janela_sistema.get_by_role("button", name="Pesquisar").first.wait_for(
                state="visible",
                timeout=15000,
            )

            return page, janela_sistema

        except Exception as erro:
            if tentativa == 0:
                print("⚠️ Falha ao navegar no menu. Acionando Clean State...")
                page = trocar_aba_aghux(
                    context,
                    page,
                    usuario_str,
                    senha_str,
                    url_aghu=url_aghu,
                )
            else:
                raise erro

    raise RuntimeError("Falha ao navegar até o módulo de Impressora por Computador.")


def _extrair_dados_linha_planilha(linha: Mapping[str, Any]) -> DadosLinhaPlanilha:
    return DadosLinhaPlanilha(
        ip_pc=str(linha["IPPC"]).strip(),
        impressora_alvo=str(linha["HostPrinter"]).strip(),
        classe_impressao=str(linha["PrinterClass"]).strip(),
    )


def _numero_linha_planilha(index: object) -> int:
    return int(str(index)) + 1


def _criar_log_linha(
    linha: Mapping[str, Any],
    dados: DadosLinhaPlanilha,
    resultado: ResultadoLinha,
) -> Row:
    return {
        "HostPC": linha.get("HostPC", ""),
        "IPPC": dados.ip_pc,
        "HostPrinter": dados.impressora_alvo,
        "IPPrinter": linha.get("IPPrinter", ""),
        "PrinterClass": dados.classe_impressao,
        "Status": resultado.status,
        "Detalhes": resultado.detalhes,
    }


def _registrar_linha_ignorada(
    index: object,
    total_linhas: int,
    campos_em_branco: list[str],
) -> ResultadoLinha:
    detalhes = (
        "Linha ignorada: campos obrigatorios em branco: "
        f"{', '.join(campos_em_branco)}."
    )

    print("\n========================================")
    print(
        f"⏭️ Ignorando linha [{_numero_linha_planilha(index)}/{total_linhas}]: "
        f"{detalhes}"
    )

    return ResultadoLinha(status="Erro", detalhes=detalhes)


def _clicar_limpar_formulario(janela_sistema) -> None:
    janela_sistema.locator("button:has(.aghu-icon-cleaner-aghu)").first.click()


def _selecionar_computador_no_formulario(
    janela_sistema,
    ip_pc: str,
    *,
    capturar_texto: bool = False,
):
    campo_computador = (
        janela_sistema.locator(
            "input[id*='computador' i], input.ui-autocomplete-input"
        )
        .locator("visible=true")
        .first
    )
    campo_computador.click()
    campo_computador.clear()
    campo_computador.press_sequentially(ip_pc, delay=150)

    padrao_exato = _criar_regex_valor_exato(ip_pc)
    caixa_flutuante_pc = (
        janela_sistema.locator("tr, li, td, span")
        .filter(has_text=padrao_exato)
        .locator("visible=true")
        .first
    )

    try:
        caixa_flutuante_pc.wait_for(state="visible", timeout=6000)
        texto_selecionado = (
            caixa_flutuante_pc.inner_text(timeout=1000) if capturar_texto else ""
        )
        caixa_flutuante_pc.click()
    except Exception as erro:
        raise ValueError("Computador não encontrado") from erro

    return campo_computador, texto_selecionado


def _buscar_computador_no_formulario(janela_sistema, ip_pc: str) -> None:
    _selecionar_computador_no_formulario(janela_sistema, ip_pc)


def _pesquisar_vinculos_computador(
    janela_sistema,
    ip_pc: str,
) -> tuple[str, list[dict]]:
    janela_sistema.get_by_role("button", name="Pesquisar").click()
    return _coletar_linhas_computador(
        janela_sistema=janela_sistema,
        ip_pc=ip_pc,
    )


def _selecionar_impressora_no_formulario(
    janela_sistema,
    impressora_alvo: str,
) -> None:
    campo_impressora = (
        janela_sistema.locator("input[id*='impressora' i]")
        .locator("visible=true")
        .first
    )
    campo_impressora.click()
    campo_impressora.clear()
    campo_impressora.press_sequentially(impressora_alvo, delay=150)

    caixa_flutuante_imp = (
        janela_sistema.locator("li, td, span")
        .filter(has_text=impressora_alvo)
        .locator("visible=true")
        .first
    )

    try:
        caixa_flutuante_imp.wait_for(state="visible", timeout=6000)
        caixa_flutuante_imp.click()
    except Exception as erro:
        raise ValueError("Impressora não existe") from erro


def _selecionar_classe_vinculo(janela_sistema, classe_impressao: str) -> None:
    campo_classe = (
        janela_sistema.locator("input[id*='classe' i], input[id*='impressao' i]")
        .locator("visible=true")
        .last
    )
    classe_atual = str(campo_classe.input_value())
    classe_aghu = _classe_impressao_aghu(classe_impressao)

    if classe_aghu.upper() in classe_atual.upper():
        return

    try:
        (
            janela_sistema.locator("button:has(.aghu-icon-cleaner-aghu)")
            .locator("visible=true")
            .last
            .click(timeout=2000)
        )
    except Exception:
        campo_classe.clear()

    botao_lupa = (
        janela_sistema.locator("button:has(.ui-icon-triangle-1-s)")
        .locator("visible=true")
        .last
    )
    botao_lupa.click()

    caixa_flutuante_classe = (
        janela_sistema.locator("li, td, span")
        .filter(has_text=classe_aghu)
        .locator("visible=true")
        .first
    )
    caixa_flutuante_classe.wait_for(state="visible", timeout=5000)
    caixa_flutuante_classe.click()


def _avaliar_pesquisa_vinculos(
    janela_sistema,
    page: Page,
    dados: DadosLinhaPlanilha,
    estado_pesquisa: str,
    registros_linhas: list[dict],
) -> tuple[ResultadoLinha | None, AcaoVinculo | None]:
    if estado_pesquisa == "indefinido":
        print("Pesquisa sem estado final claro. Conferir manualmente.")
        _limpar_estado_formulario(janela_sistema, page)
        return (
            ResultadoLinha(
                status="Erro",
                detalhes=MENSAGEM_ERRO_PESQUISA_INDEFINIDA,
            ),
            None,
        )

    if estado_pesquisa == "linhas" and not registros_linhas:
        print("Pesquisa retornou linhas, mas nenhuma com o IP esperado.")
        _limpar_estado_formulario(janela_sistema, page)
        return (
            ResultadoLinha(
                status="Erro",
                detalhes=(
                    "Conferir manualmente: pesquisa retornou linhas, mas nenhuma "
                    f"com o IP esperado [{dados.ip_pc}]."
                ),
            ),
            None,
        )

    decisao_linha, registro_linha = _decidir_acao_linhas(
        registros_linhas=registros_linhas,
        impressora_alvo=dados.impressora_alvo,
        classe_impressao=dados.classe_impressao,
    )

    if decisao_linha != "conferir":
        return None, AcaoVinculo(tipo=decisao_linha, registro=registro_linha)

    tipo_cups_atual = (
        str(registro_linha.get("tipo_cups") or "nao identificado")
        if registro_linha
        else "nao identificado"
    )
    print("Tipo do Cups divergente. Conferir manualmente.")
    _clicar_limpar_formulario(janela_sistema)

    return (
        ResultadoLinha(
            status="Erro",
            detalhes=(
                "Conferir manualmente: impressora ja vinculada ao computador "
                f"com Tipo do Cups [{tipo_cups_atual}], diferente da planilha "
                f"[{dados.classe_impressao}]."
            ),
        ),
        None,
    )


def _gravar_formulario_vinculo(
    janela_sistema,
    page: Page,
    *,
    operacao: str,
) -> ResultadoLinha | None:
    janela_sistema.get_by_role("button", name="Gravar").click()
    resultado_gravacao, mensagem_gravacao = _aguardar_resultado_gravacao(
        janela_sistema,
        page,
    )

    if resultado_gravacao == "erro":
        print(f"Erro retornado pelo AGHU: {mensagem_gravacao}")

        if operacao == "incluir" and _erro_classe_pdf_duplicada(mensagem_gravacao):
            detalhes = (
                "Conferir manualmente: AGHU bloqueou inclusao porque "
                "ja existe impressora na classe A/PDF para o computador."
            )
        else:
            detalhes = (
                f"Conferir manualmente: AGHU retornou erro ao {operacao}: "
                f"{mensagem_gravacao}"
            )

        _limpar_estado_formulario(janela_sistema, page)
        return ResultadoLinha(status="Erro", detalhes=detalhes)

    if resultado_gravacao == "indefinido":
        print("Gravacao sem sucesso ou erro conhecido. Conferir manualmente.")
        _limpar_estado_formulario(janela_sistema, page)
        return ResultadoLinha(
            status="Erro",
            detalhes=(
                "Conferir manualmente: gravacao nao retornou sucesso "
                "nem erro conhecido."
            ),
        )

    return None


def _registrar_vinculo_mantido(janela_sistema) -> ResultadoLinha:
    print("✅ SUCESSO! A impressora já estava correta.")
    _clicar_limpar_formulario(janela_sistema)
    return ResultadoLinha(
        status="Mantido",
        detalhes="Impressora já estava correta no sistema.",
    )


def _linha_alvo_da_acao(acao: AcaoVinculo):
    linha_alvo = acao.registro["linha"] if acao.registro else None

    if linha_alvo is None:
        raise RuntimeError("Linha da tabela para decisao nao localizada.")

    return linha_alvo


def _editar_vinculo_existente(
    janela_sistema,
    page: Page,
    linha_alvo,
    dados: DadosLinhaPlanilha,
    impressora_fabricada_agora: bool,
) -> ResultadoLinha:
    print("⚠️ DIVERGÊNCIA! Atualizando a impressora...")
    botao_lapis = linha_alvo.locator(
        'td.first-column.auto-adjust a[title="editar"], '
        '[title*="editar" i], [title*="alterar" i], .aghu-icon-edit'
    ).first
    botao_lapis.click()
    janela_sistema.get_by_role("button", name="Gravar").wait_for(state="visible")

    janela_sistema.locator("button:has(.aghu-icon-cleaner-aghu)").click()
    _selecionar_impressora_no_formulario(janela_sistema, dados.impressora_alvo)

    erro_gravacao = _gravar_formulario_vinculo(
        janela_sistema,
        page,
        operacao="alterar",
    )

    if erro_gravacao is not None:
        return erro_gravacao

    _aguardar_botao_pesquisar_se_possivel(janela_sistema)
    print("🔄 Salvamento concluído!")

    _clicar_limpar_formulario(janela_sistema)

    if impressora_fabricada_agora:
        return ResultadoLinha(
            status="Criado",
            detalhes="Impressora cadastrada no CUPS e atualizada.",
        )

    return ResultadoLinha(
        status="Alterado",
        detalhes="Vínculo atualizado com sucesso.",
    )


def _incluir_novo_vinculo(
    janela_sistema,
    page: Page,
    dados: DadosLinhaPlanilha,
    registros_linhas: list[dict],
    impressora_fabricada_agora: bool,
) -> ResultadoLinha:
    if registros_linhas:
        print("Vinculos existentes nao sao PDF. Iniciando NOVO vinculo...")
    else:
        print("Nenhum registro encontrado. Iniciando NOVO vinculo...")

    janela_sistema.get_by_role("button", name="Novo").click()
    janela_sistema.get_by_role("button", name="Gravar").wait_for(state="visible")

    campo_computador, texto_computador_selecionado = (
        _selecionar_computador_no_formulario(
            janela_sistema,
            dados.ip_pc,
            capturar_texto=True,
        )
    )
    computador_confere, computador_selecionado = _validar_computador_selecionado(
        campo_computador,
        dados.ip_pc,
        texto_computador_selecionado,
    )

    if not computador_confere:
        raise ValueError(
            "Conferir manualmente: computador selecionado diverge "
            f"do IP esperado. Esperado: {dados.ip_pc}; selecionado: "
            f"{computador_selecionado}."
        )

    _selecionar_impressora_no_formulario(janela_sistema, dados.impressora_alvo)
    _selecionar_classe_vinculo(janela_sistema, dados.classe_impressao)

    erro_gravacao = _gravar_formulario_vinculo(
        janela_sistema,
        page,
        operacao="incluir",
    )

    if erro_gravacao is not None:
        return erro_gravacao

    _aguardar_botao_pesquisar_se_possivel(janela_sistema)
    print("🔄 Cadastro finalizado com sucesso!")

    _clicar_limpar_formulario(janela_sistema)

    if impressora_fabricada_agora:
        return ResultadoLinha(
            status="Criado",
            detalhes="Impressora nova identificada no CUPS, criada e vinculada no AGHU.",
        )

    return ResultadoLinha(
        status="Vinculado",
        detalhes="Vínculo criado com sucesso.",
    )


def _executar_acao_vinculo(
    janela_sistema,
    page: Page,
    dados: DadosLinhaPlanilha,
    registros_linhas: list[dict],
    acao: AcaoVinculo,
    impressora_fabricada_agora: bool,
) -> ResultadoLinha:
    if acao.tipo == "mantido":
        return _registrar_vinculo_mantido(janela_sistema)

    if acao.tipo == "alterar":
        return _editar_vinculo_existente(
            janela_sistema=janela_sistema,
            page=page,
            linha_alvo=_linha_alvo_da_acao(acao),
            dados=dados,
            impressora_fabricada_agora=impressora_fabricada_agora,
        )

    return _incluir_novo_vinculo(
        janela_sistema=janela_sistema,
        page=page,
        dados=dados,
        registros_linhas=registros_linhas,
        impressora_fabricada_agora=impressora_fabricada_agora,
    )


def _processar_linha_aghu(
    janela_sistema,
    page: Page,
    dados: DadosLinhaPlanilha,
    impressora_fabricada_agora: bool,
) -> ResultadoLinha:
    passo_atual = "Iniciando"

    try:
        passo_atual = "Buscando Computador"
        _buscar_computador_no_formulario(janela_sistema, dados.ip_pc)

        passo_atual = "Pesquisando na Tabela"
        estado_pesquisa, registros_linhas = _pesquisar_vinculos_computador(
            janela_sistema=janela_sistema,
            ip_pc=dados.ip_pc,
        )
        resultado, acao = _avaliar_pesquisa_vinculos(
            janela_sistema=janela_sistema,
            page=page,
            dados=dados,
            estado_pesquisa=estado_pesquisa,
            registros_linhas=registros_linhas,
        )

        if resultado is not None:
            return resultado

        if acao is None:
            return ResultadoLinha(status="Erro", detalhes="Falha Desconhecida.")

        if acao.tipo == "alterar":
            passo_atual = "Editando Impressora Existente"
        elif acao.tipo == "incluir":
            passo_atual = "Cadastrando Nova Impressora (Vinculando)"

        return _executar_acao_vinculo(
            janela_sistema=janela_sistema,
            page=page,
            dados=dados,
            registros_linhas=registros_linhas,
            acao=acao,
            impressora_fabricada_agora=impressora_fabricada_agora,
        )

    except ValueError:
        raise
    except Exception as erro:
        raise FalhaTecnicaProcessamento(passo_atual, erro) from erro


def _reiniciar_modulo_vinculo(
    context: BrowserContext,
    page: Page,
    usuario_str: str,
    senha_str: str,
    url_aghu: str,
) -> tuple[Page, object]:
    page = trocar_aba_aghux(
        context,
        page,
        usuario_str,
        senha_str,
        url_aghu=url_aghu,
    )
    return navegar_ate_modulo(
        context,
        page,
        usuario_str,
        senha_str,
        url_aghu=url_aghu,
    )


def _cadastrar_impressora_inexistente(
    context: BrowserContext,
    page: Page,
    usuario_str: str,
    senha_str: str,
    dados: DadosLinhaPlanilha,
    url_aghu: str,
) -> tuple[bool, str, Page]:
    erro_estoquista = ""

    for tentativa_estoque in range(MAX_TENTATIVAS_ESTOQUE):
        try:
            print(
                "👷 [Estoquista - Tentativa "
                f"{tentativa_estoque + 1}/{MAX_TENTATIVAS_ESTOQUE}] "
                "Isolando ambiente..."
            )
            page = trocar_aba_aghux(
                context,
                page,
                usuario_str,
                senha_str,
                url_aghu=url_aghu,
            )

            dados_cups = consultar_dados_site_secundario(
                context,
                dados.impressora_alvo,
                dados.classe_impressao,
            )
            janela_cadastro_imp = navegar_ate_cadastro_impressora(page)
            cadastrar_nova_impressora(janela_cadastro_imp, dados_cups)

            return True, "", page

        except Exception as e_cups:
            erro_estoquista = str(e_cups)
            print(f"⚠️ O Estoquista tropeçou: {erro_estoquista}")

    return False, erro_estoquista, page


def _tratar_impressora_inexistente(
    context: BrowserContext,
    page: Page,
    janela_sistema,
    usuario_str: str,
    senha_str: str,
    dados: DadosLinhaPlanilha,
    url_aghu: str,
) -> tuple[bool, ResultadoLinha, Page, object]:
    print("🚨 ALERTA: A impressora não está no AGHUX! Chamando o Robô Estoquista...")
    sucesso_estoquista, erro_estoquista, page = _cadastrar_impressora_inexistente(
        context=context,
        page=page,
        usuario_str=usuario_str,
        senha_str=senha_str,
        dados=dados,
        url_aghu=url_aghu,
    )

    if sucesso_estoquista:
        print(
            "🔙 O Estoquista terminou! Maestro criando nova aba limpa para "
            "retomar o vínculo..."
        )
        page, janela_sistema = _reiniciar_modulo_vinculo(
            context=context,
            page=page,
            usuario_str=usuario_str,
            senha_str=senha_str,
            url_aghu=url_aghu,
        )
        print("🔄 Estoque abastecido! Gastando uma Vida do Vinculador para tentar de novo...")
        return (
            True,
            ResultadoLinha(status="Erro", detalhes="Falha Desconhecida."),
            page,
            janela_sistema,
        )

    if "Não existe no CUPS" in erro_estoquista:
        resultado = ResultadoLinha(
            status="Inexistente",
            detalhes="Fila de impressão não encontrada no Servidor CUPS.",
        )
    else:
        resultado = ResultadoLinha(
            status="Erro",
            detalhes=f"Falha no Almoxarifado após 3 tentativas: {erro_estoquista}",
        )

    print(
        "❌ O Estoquista falhou definitivamente: "
        f"{resultado.status} - {resultado.detalhes}"
    )
    return False, resultado, page, janela_sistema


def _limpar_apos_erro_funcional(janela_sistema) -> None:
    try:
        janela_sistema.get_by_role("button", name="Cancelar").click(timeout=1000)
    except Exception:
        pass

    try:
        janela_sistema.locator(
            "button:has(.aghu-icon-cleaner-aghu)"
        ).first.click(timeout=1500)
    except Exception:
        pass


def _tratar_erro_funcional(
    janela_sistema,
    mensagem_erro: str,
) -> ResultadoLinha:
    if "Computador não encontrado" in mensagem_erro:
        resultado = ResultadoLinha(
            status="Inexistente",
            detalhes="Computador não cadastrado no AGHUX.",
        )
    else:
        resultado = ResultadoLinha(status="Erro", detalhes=mensagem_erro)

    print(f"❌ Identificado erro sem salvação imediata: {resultado.status}")
    _limpar_apos_erro_funcional(janela_sistema)
    return resultado


def _tratar_falha_tecnica(
    context: BrowserContext,
    page: Page,
    janela_sistema,
    usuario_str: str,
    senha_str: str,
    url_aghu: str,
    tentativa: int,
    passo_atual: str,
) -> tuple[ResultadoLinha | None, Page, object]:
    print(f"❌ O navegador congelou ou não achou o elemento no passo: '{passo_atual}'")

    if tentativa < MAX_TENTATIVAS_PROCESSAMENTO_LINHA - 1:
        print(
            "⚡ Pegando o Desfibrilador! Iniciando tentativa "
            f"{tentativa + 2}/{MAX_TENTATIVAS_PROCESSAMENTO_LINHA} em nova aba..."
        )
        try:
            page, janela_sistema = _reiniciar_modulo_vinculo(
                context=context,
                page=page,
                usuario_str=usuario_str,
                senha_str=senha_str,
                url_aghu=url_aghu,
            )
            print("🔄 Sistema ressuscitado em nova aba limpa. Retomando a missão!")
        except Exception as e_recup:
            print(f"⚠️ A ressuscitação falhou: {e_recup}")

        return None, page, janela_sistema

    resultado = ResultadoLinha(
        status="Erro",
        detalhes=(
            f"Falha de sistema ou rede no passo '{passo_atual}' após "
            f"{MAX_TENTATIVAS_PROCESSAMENTO_LINHA} tentativas."
        ),
    )
    print("🚨 As 3 vidas acabaram. O sistema está instável.")

    try:
        page, janela_sistema = _reiniciar_modulo_vinculo(
            context=context,
            page=page,
            usuario_str=usuario_str,
            senha_str=senha_str,
            url_aghu=url_aghu,
        )
    except Exception:
        pass

    return resultado, page, janela_sistema


def _processar_linha_com_retentativas(
    context: BrowserContext,
    page: Page,
    janela_sistema,
    usuario_str: str,
    senha_str: str,
    dados: DadosLinhaPlanilha,
    url_aghu: str,
) -> tuple[ResultadoLinha, Page, object]:
    resultado = ResultadoLinha(status="Erro", detalhes="Falha Desconhecida.")
    impressora_fabricada_agora = False

    for tentativa in range(MAX_TENTATIVAS_PROCESSAMENTO_LINHA):
        try:
            resultado = _processar_linha_aghu(
                janela_sistema=janela_sistema,
                page=page,
                dados=dados,
                impressora_fabricada_agora=impressora_fabricada_agora,
            )
            return resultado, page, janela_sistema

        except ValueError as erro:
            mensagem_erro = str(erro)

            if mensagem_erro == "Impressora não existe":
                (
                    tentar_novamente,
                    resultado,
                    page,
                    janela_sistema,
                ) = _tratar_impressora_inexistente(
                    context=context,
                    page=page,
                    janela_sistema=janela_sistema,
                    usuario_str=usuario_str,
                    senha_str=senha_str,
                    dados=dados,
                    url_aghu=url_aghu,
                )

                if tentar_novamente:
                    impressora_fabricada_agora = True
                    continue

                return resultado, page, janela_sistema

            resultado = _tratar_erro_funcional(janela_sistema, mensagem_erro)
            return resultado, page, janela_sistema

        except FalhaTecnicaProcessamento as erro:
            resultado_tecnico, page, janela_sistema = _tratar_falha_tecnica(
                context=context,
                page=page,
                janela_sistema=janela_sistema,
                usuario_str=usuario_str,
                senha_str=senha_str,
                url_aghu=url_aghu,
                tentativa=tentativa,
                passo_atual=erro.passo,
            )

            if resultado_tecnico is not None:
                return resultado_tecnico, page, janela_sistema

    return resultado, page, janela_sistema


def _gerar_relatorio_xlsx(
    logs_do_diario: Sequence[Mapping[str, Any]],
    source_spreadsheet_path: str,
    report_path: Path,
) -> str:
    print("\nFim da leitura! Gerando relatorio XLSX...")
    write_report(source_spreadsheet_path, report_path, logs_do_diario)
    print(f"Relatorio gerado com sucesso: {report_path}")
    return str(report_path)

# ==========================================
# CAPÍTULO 3: O CÉREBRO MAESTRO
# ==========================================
def processar_computadores(
    context: BrowserContext,
    page_inicial: Page,
    janela_sistema_inicial,
    caminho_planilha: str,
    usuario_str: str,
    senha_str: str,
    report_directory: str | os.PathLike,
    url_aghu: str = AGHU_URL,
) -> str:
    logs_do_diario = []
    page = page_inicial
    janela_sistema = janela_sistema_inicial
    planilha = ler_planilha(caminho_planilha)
    report_path = build_report_path(str(report_directory), caminho_planilha)
    total_linhas = len(planilha)

    for index, linha in enumerate(planilha):
        dados = _extrair_dados_linha_planilha(linha)
        campos_em_branco = _campos_obrigatorios_planilha_em_branco(linha)

        if campos_em_branco:
            resultado = _registrar_linha_ignorada(
                index=index,
                total_linhas=total_linhas,
                campos_em_branco=campos_em_branco,
            )
            logs_do_diario.append(_criar_log_linha(linha, dados, resultado))
            continue

        print("\n========================================")
        print(
            f"🔍 Investigando [{_numero_linha_planilha(index)}/{total_linhas}]: "
            f"Computador [{dados.ip_pc}] | Alvo [{dados.impressora_alvo}]"
        )

        resultado, page, janela_sistema = _processar_linha_com_retentativas(
            context=context,
            page=page,
            janela_sistema=janela_sistema,
            usuario_str=usuario_str,
            senha_str=senha_str,
            dados=dados,
            url_aghu=url_aghu,
        )

        print(f"📝 Anotando no diário: [{resultado.status}] {resultado.detalhes}")
        logs_do_diario.append(_criar_log_linha(linha, dados, resultado))

    return _gerar_relatorio_xlsx(
        logs_do_diario=logs_do_diario,
        source_spreadsheet_path=caminho_planilha,
        report_path=report_path,
    )
