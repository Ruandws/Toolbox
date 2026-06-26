from pathlib import Path

import pytest

import cadastro_pessoa_aghu as aghu
from cadastro_pessoa_aghu import (
    CadastroPessoaEntrada,
    FluxoResultado,
    ResultadoCadastroPessoa,
    STATUS_ATUALIZADO,
    STATUS_CONFERIR_MANUAL,
    STATUS_CRIADO,
    STATUS_ERRO,
    STATUS_IGNORADO,
    executar_cadastro_individual,
    executar_cadastro_lote,
    executar_cadastro_pessoas,
    ler_planilha_cadastros,
    salvar_relatorio_resultados,
)
from testes.aghu.fixtures_cadastro_pessoa import FakeContext, pessoa_valida


class TestNormalizacao:
    def test_valor_em_branco_reconhece_vazios(self):
        import pandas as pd

        assert aghu._valor_em_branco("") is True
        assert aghu._valor_em_branco("   ") is True
        assert aghu._valor_em_branco(None) is True
        assert aghu._valor_em_branco(pd.NA) is True
        assert aghu._valor_em_branco(float("nan")) is True
        assert aghu._valor_em_branco("abc") is False

    def test_normalizar_texto_colapsa_espacos_e_casefold(self):
        assert aghu.normalizar_texto("  A\tB\nC  ") == "a b c"

    def test_texto_planilha_remove_espacos_e_nan(self):
        assert aghu.texto_planilha("  valor  ") == "valor"
        assert aghu.texto_planilha(None) == ""

    def test_apenas_digitos_remove_mascara(self):
        assert aghu.apenas_digitos("CPF 123.456.789-01") == "12345678901"

    def test_cpf_confere_por_digitos_finais(self):
        assert aghu.cpf_confere("123.456.789-01", "45678901") is True
        assert aghu.cpf_confere("123.456.789-01", "00000000000") is False

    def test_normalizar_entrada_limpa_campos_e_cpf(self):
        entrada = pessoa_valida(
            nome_pessoa="  Joao   Silva  ",
            cpf="  123.456.789-01  ",
        )

        normalizada = aghu.normalizar_entrada(entrada)

        assert normalizada.nome_pessoa == "Joao   Silva"
        assert normalizada.cpf == "12345678901"


class TestValidacao:
    def test_entrada_valida_nao_retorna_erros(self):
        assert aghu.validar_entrada(pessoa_valida()) == []

    def test_campos_obrigatorios_em_branco_retorna_erro_unico_com_colunas(self):
        erros = aghu.validar_entrada(CadastroPessoaEntrada())

        assert len(erros) == 1
        assert "campos obrigatorios em branco" in erros[0]
        assert aghu._nome_coluna("nome_pessoa") in erros[0]
        assert aghu._nome_coluna("cpf") in erros[0]

    def test_cpf_com_tamanho_invalido_retorna_erro(self):
        erros = aghu.validar_entrada(pessoa_valida(cpf="123"))

        assert "CPF deve conter 11 digitos" in erros[0]

    def test_resultado_ignorado_normaliza_entrada(self):
        resultado = aghu.resultado_ignorado(
            pessoa_valida(nome_pessoa=" Ana ", cpf="123.456.789-01"),
            "motivo",
        )

        assert resultado.status == STATUS_IGNORADO
        assert resultado.nome_pessoa == "Ana"
        assert resultado.cpf == "12345678901"
        assert resultado.detalhes == "motivo"


class TestDataclasses:
    def test_cadastro_pessoa_entrada_e_frozen(self):
        entrada = pessoa_valida()

        with pytest.raises(AttributeError):
            entrada.cpf = "outro"  # type: ignore[misc]

    def test_resultado_cadastro_pessoa_e_frozen(self):
        resultado = ResultadoCadastroPessoa(
            cpf="123",
            nome_pessoa="Nome",
            status=STATUS_ERRO,
            detalhes="erro",
        )

        with pytest.raises(AttributeError):
            resultado.status = STATUS_CRIADO  # type: ignore[misc]


class TestPlanilha:
    def test_ler_planilha_valida_com_colunas_principais(self, tmp_xlsx):
        campos = list(aghu.ALIASES_COLUNAS)
        cabecalhos = [aghu.ALIASES_COLUNAS[campo][0] for campo in campos]
        linha = [getattr(pessoa_valida(), campo) for campo in campos]
        caminho = tmp_xlsx("pessoas.xlsx", cabecalhos, [linha])

        cadastros = ler_planilha_cadastros(caminho)

        assert len(cadastros) == 1
        assert cadastros[0].nome_pessoa == "Joao Silva"
        assert cadastros[0].cpf == "12345678901"

    def test_ler_planilha_reconhece_aliases_sem_acentos(self, tmp_xlsx):
        caminho = tmp_xlsx(
            "aliases.xlsx",
            [
                "Nome",
                "Nome Mae",
                "Sexo",
                "Nascimento",
                "Nacionalidade",
                "Naturalidade",
                "RG",
                "Orgao Emissor",
                "UF RG",
                "CPF",
            ],
            [
                [
                    "Ana",
                    "Maria",
                    "Feminino",
                    "02/02/1991",
                    "Brasileira",
                    "Brasilia",
                    "123",
                    "SSP",
                    "DF",
                    "987.654.321-00",
                ]
            ],
        )

        cadastros = ler_planilha_cadastros(caminho)

        assert cadastros[0].nome_mae == "Maria"
        assert cadastros[0].rg == "123"
        assert cadastros[0].orgao_emissor == "SSP"
        assert cadastros[0].cpf == "98765432100"

    def test_ler_planilha_remove_espacos_dos_cabecalhos(self, tmp_xlsx):
        campos = list(aghu.CAMPOS_PESSOA_OBRIGATORIOS)
        cabecalhos = [f" {aghu.ALIASES_COLUNAS[campo][0]} " for campo in campos]
        linha = [getattr(pessoa_valida(), campo) for campo in campos]
        caminho = tmp_xlsx("cabecalhos.xlsx", cabecalhos, [linha])

        cadastros = ler_planilha_cadastros(caminho)

        assert cadastros[0].cpf == "12345678901"

    def test_ler_planilha_inexistente_levanta_erro(self):
        with pytest.raises(FileNotFoundError):
            ler_planilha_cadastros("nao_existe.xlsx")

    def test_ler_planilha_rejeita_extensao_invalida(self, tmp_path):
        caminho = tmp_path / "pessoas.csv"
        caminho.write_text("x")

        with pytest.raises(ValueError, match=".xlsx"):
            ler_planilha_cadastros(caminho)

    def test_ler_planilha_rejeita_colunas_obrigatorias_ausentes(self, tmp_xlsx):
        caminho = tmp_xlsx("incompleta.xlsx", ["CPF"], [["12345678901"]])

        with pytest.raises(ValueError, match="Colunas obrigatorias ausentes"):
            ler_planilha_cadastros(caminho)


class TestRelatoriosLogs:
    def test_salvar_relatorio_cria_xlsx_com_colunas_esperadas(self, tmp_path):
        import openpyxl

        resultado = ResultadoCadastroPessoa(
            cpf="12345678901",
            nome_pessoa="Joao",
            status=STATUS_CRIADO,
            detalhes="OK",
            pessoa="criado: OK",
        )

        caminho = salvar_relatorio_resultados(
            [resultado],
            tmp_path / "saida_sem_extensao",
        )

        wb = openpyxl.load_workbook(caminho)
        ws = wb["Resultado"]
        cabecalhos = [celula.value for celula in ws[1]]
        valores = [celula.value for celula in ws[2]]
        wb.close()

        assert caminho.suffix == ".xlsx"
        assert cabecalhos == ["CPF", "Nome da Pessoa", "Status", "Detalhes", "Pessoa"]
        assert valores == ["12345678901", "Joao", STATUS_CRIADO, "OK", "criado: OK"]

    def test_salvar_relatorio_rejeita_extensao_invalida(self, tmp_path):
        with pytest.raises(ValueError, match=".xlsx"):
            salvar_relatorio_resultados([], tmp_path / "saida.csv")

    def test_gerar_csv_logs_cria_arquivo_com_usuario_e_resultados(self, tmp_path):
        resultado = ResultadoCadastroPessoa(
            cpf="12345678901",
            nome_pessoa="Joao",
            status=STATUS_ERRO,
            detalhes="falha",
        )

        caminho = Path(aghu.gerar_csv_logs([resultado], "usuario.rede", tmp_path))
        texto = caminho.read_text(encoding="utf-8-sig")

        assert caminho.exists()
        assert "Atualizado por: usuario.rede" in texto
        assert "12345678901;Joao;erro;falha" in texto

    def test_controle_saida_terminal_suprime_stdout_e_stderr(self, monkeypatch, capsys):
        import sys

        monkeypatch.setattr(aghu, "esconder_console_windows", lambda: None)

        with aghu.controle_saida_terminal(False):
            print("stdout oculto")
            print("stderr oculto", file=sys.stderr)

        saida = capsys.readouterr()

        assert saida.out == ""
        assert saida.err == ""

    def test_controle_saida_terminal_mantem_saida_quando_console_ativo(self, capsys):
        with aghu.controle_saida_terminal(True):
            print("stdout visivel")

        assert "stdout visivel" in capsys.readouterr().out


class FakeCell:
    def __init__(self, texto):
        self.texto = texto

    def inner_text(self, timeout=500):
        return self.texto


class FakeCells:
    def __init__(self, textos):
        self.textos = textos

    def count(self):
        return len(self.textos)

    def nth(self, indice):
        return FakeCell(self.textos[indice])


class FakeRow:
    def __init__(self, textos=None, classe="", texto_linha="", visivel=True):
        self.textos = textos or []
        self.classe = classe
        self.texto_linha = texto_linha
        self.visivel = visivel

    def get_attribute(self, nome, timeout=250):
        assert nome == "class"
        return self.classe

    def locator(self, seletor):
        assert seletor == "td"
        return FakeCells(self.textos)

    def is_visible(self, timeout=250):
        return self.visivel

    def inner_text(self, timeout=250):
        return self.texto_linha


class FakeRows:
    def __init__(self, linhas):
        self.linhas = linhas

    def count(self):
        return len(self.linhas)

    def nth(self, indice):
        return self.linhas[indice]

    @property
    def first(self):
        return self.nth(0)


class TestLocatorsPuros:
    def test_linha_vazia_visivel_por_classe(self):
        linhas = FakeRows([FakeRow(classe="ui-datatable-empty-message")])

        assert aghu.linha_vazia_visivel(linhas) is True

    def test_linha_vazia_visivel_por_texto(self):
        linhas = FakeRows([FakeRow(texto_linha=aghu.TEXTO_NENHUM_REGISTRO)])

        assert aghu.linha_vazia_visivel(linhas) is True

    def test_linha_vazia_ignora_linhas_invisiveis(self):
        linhas = FakeRows(
            [
                FakeRow(
                    classe="ui-datatable-empty-message",
                    texto_linha=aghu.TEXTO_NENHUM_REGISTRO,
                    visivel=False,
                )
            ]
        )

        assert aghu.linha_vazia_visivel(linhas) is False

    def test_linha_por_cpf_retorna_linha_com_cpf_exato(self):
        linha = FakeRow(["", "", "", "", "123.456.789-01"])
        linhas = FakeRows([FakeRow(classe="ui-datatable-empty-message"), linha])
        flow = aghu.PessoaFlow(object())

        assert flow._linha_por_cpf(linhas, "12345678901") is linha

    def test_linha_por_cpf_ignora_linhas_sem_colunas_suficientes(self):
        linhas = FakeRows([FakeRow(["1", "2"]), FakeRow(["", "", "", "", "000"])])
        flow = aghu.PessoaFlow(object())

        assert flow._linha_por_cpf(linhas, "12345678901") is None


class FakeMessageItem:
    def __init__(self, texto, visivel=True):
        self.texto = texto
        self.visivel = visivel

    def is_visible(self, timeout=250):
        return self.visivel

    def inner_text(self, timeout=500):
        return self.texto


class FakeMessageCollection:
    def __init__(self, itens):
        self.itens = itens

    def count(self):
        return len(self.itens)

    def nth(self, indice):
        return self.itens[indice]


class FakeJanelaMensagens:
    def __init__(self, mensagens):
        self.mensagens = mensagens

    def locator(self, _seletor):
        return FakeMessageCollection(self.mensagens)


class TestMensagens:
    def test_mensagens_sistema_retorna_apenas_visiveis_com_texto(self):
        janela = FakeJanelaMensagens(
            [
                FakeMessageItem(" OK "),
                FakeMessageItem("oculta", visivel=False),
                FakeMessageItem(""),
            ]
        )

        assert aghu.mensagens_sistema(janela) == ["OK"]

    def test_aguardar_mensagem_gravacao_retorna_sucesso(self):
        janela = FakeJanelaMensagens([FakeMessageItem("Registro gravado")])

        status, mensagem = aghu.aguardar_mensagem_gravacao(
            janela,
            sucessos=("gravado",),
        )

        assert status == "sucesso"
        assert mensagem == "Registro gravado"

    def test_aguardar_mensagem_gravacao_retorna_erro_negocio(self):
        janela = FakeJanelaMensagens([FakeMessageItem("CPF duplicado")])

        status, mensagem = aghu.aguardar_mensagem_gravacao(
            janela,
            sucessos=("gravado",),
            erros_negocio=("duplicado",),
        )

        assert status == "erro_negocio"
        assert mensagem == "CPF duplicado"

    def test_aguardar_mensagem_gravacao_retorna_erro_generico(self):
        janela = FakeJanelaMensagens([FakeMessageItem("Campo obrigatorio ausente")])

        status, mensagem = aghu.aguardar_mensagem_gravacao(
            janela,
            sucessos=("gravado",),
        )

        assert status == "erro"
        assert mensagem == "Campo obrigatorio ausente"

    def test_aguardar_mensagem_gravacao_indefinida_no_timeout(self):
        janela = FakeJanelaMensagens([])

        status, mensagem = aghu.aguardar_mensagem_gravacao(
            janela,
            sucessos=("gravado",),
            timeout_ms=0,
        )

        assert status == "indefinido"
        assert "tempo limite" in mensagem


class TestPessoaFlow:
    def test_normalizar_sexo(self):
        assert aghu.PessoaFlow._normalizar_sexo("Masculino") == "Masculino"
        assert aghu.PessoaFlow._normalizar_sexo("feminino") == "Feminino"
        assert aghu.PessoaFlow._normalizar_sexo("outro") == "Ignorado"

    def test_processar_atualiza_quando_cpf_encontrado(self, monkeypatch):
        flow = aghu.PessoaFlow(object())
        linha = object()
        chamadas = []
        monkeypatch.setattr(flow, "pesquisar_por_cpf", lambda cpf: ("encontrado", linha))
        monkeypatch.setattr(aghu, "clicar_acao", lambda alvo, nomes: chamadas.append((alvo, nomes)))
        monkeypatch.setattr(flow, "_preencher_formulario", lambda entrada: chamadas.append(("preencher", entrada.cpf)))
        monkeypatch.setattr(flow, "_gravar_pessoa", lambda: ("sucesso", "alterada"))
        monkeypatch.setattr(flow, "_retornar_para_pesquisa", lambda: chamadas.append("voltar"))

        resultado = flow.processar(pessoa_valida())

        assert resultado == FluxoResultado(STATUS_ATUALIZADO, "alterada")
        assert (linha, ("Editar",)) in chamadas
        assert "voltar" in chamadas

    def test_processar_cria_quando_cpf_nao_encontrado(self, monkeypatch):
        flow = aghu.PessoaFlow(object())
        chamadas = []
        monkeypatch.setattr(flow, "pesquisar_por_cpf", lambda cpf: ("nao_encontrado", None))
        monkeypatch.setattr(aghu, "clicar_botao", lambda janela, nome: chamadas.append(nome))
        monkeypatch.setattr(aghu, "aguardar_ciclo_carregamento", lambda janela, deteccao_ms=500: True)
        monkeypatch.setattr(flow, "_preencher_formulario", lambda entrada: chamadas.append(("preencher", entrada.cpf)))
        monkeypatch.setattr(flow, "_gravar_pessoa", lambda: ("sucesso", "incluida"))
        monkeypatch.setattr(flow, "_retornar_para_pesquisa", lambda: chamadas.append("voltar"))

        resultado = flow.processar(pessoa_valida())

        assert resultado == FluxoResultado(STATUS_CRIADO, "incluida")
        assert "Novo" in chamadas
        assert "voltar" in chamadas

    def test_processar_retorna_conferir_manual_quando_pesquisa_indefinida(self, monkeypatch):
        flow = aghu.PessoaFlow(object())
        monkeypatch.setattr(flow, "pesquisar_por_cpf", lambda cpf: ("indefinido", None))

        resultado = flow.processar(pessoa_valida())

        assert resultado.status == STATUS_CONFERIR_MANUAL
        assert "nao retornou estado conclusivo" in resultado.detalhes

    def test_processar_retorna_erro_quando_gravacao_falha(self, monkeypatch):
        flow = aghu.PessoaFlow(object())
        monkeypatch.setattr(flow, "pesquisar_por_cpf", lambda cpf: ("nao_encontrado", None))
        monkeypatch.setattr(aghu, "clicar_botao", lambda *_args, **_kwargs: None)
        monkeypatch.setattr(aghu, "aguardar_ciclo_carregamento", lambda *_args, **_kwargs: True)
        monkeypatch.setattr(flow, "_preencher_formulario", lambda entrada: None)
        monkeypatch.setattr(flow, "_gravar_pessoa", lambda: ("erro", "falha"))
        monkeypatch.setattr(flow, "_retornar_para_pesquisa", lambda: None)

        resultado = flow.processar(pessoa_valida())

        assert resultado == FluxoResultado(STATUS_ERRO, "falha")


class TestMaestro:
    def test_processar_cadastro_normaliza_entrada_e_resume_fluxo(self, monkeypatch):
        class FlowFake:
            def __init__(self, janela):
                self.janela = janela

            def processar(self, entrada):
                assert entrada.cpf == "12345678901"
                assert entrada.nome_pessoa == "Joao"
                return FluxoResultado(STATUS_CRIADO, "OK")

        monkeypatch.setattr(aghu, "PessoaFlow", FlowFake)

        resultado = aghu.processar_cadastro(
            object(),
            pessoa_valida(nome_pessoa=" Joao ", cpf="123.456.789-01"),
        )

        assert resultado == ResultadoCadastroPessoa(
            cpf="12345678901",
            nome_pessoa="Joao",
            status=STATUS_CRIADO,
            detalhes="OK",
            pessoa="criado: OK",
        )

    def test_processar_cadastros_ignora_linha_invalida_sem_browser(self, monkeypatch):
        monkeypatch.setattr(
            aghu,
            "garantir_tela_pesquisa_pessoa",
            lambda *_args, **_kwargs: pytest.fail("nao deveria navegar"),
        )

        resultados = aghu.processar_cadastros(
            context=object(),
            page_inicial=object(),
            janela_sistema_inicial=object(),
            cadastros=[CadastroPessoaEntrada(cpf="123")],
            usuario_rede="usuario",
            senha="senha",
        )

        assert len(resultados) == 1
        assert resultados[0].status == STATUS_IGNORADO
        assert "Linha ignorada" in resultados[0].detalhes

    def test_processar_cadastros_processa_linha_valida(self, monkeypatch):
        monkeypatch.setattr(
            aghu,
            "garantir_tela_pesquisa_pessoa",
            lambda **kwargs: (kwargs["page_atual"], kwargs["janela_atual"]),
        )
        monkeypatch.setattr(
            aghu,
            "processar_cadastro",
            lambda janela, entrada: ResultadoCadastroPessoa(
                cpf=entrada.cpf,
                nome_pessoa=entrada.nome_pessoa,
                status=STATUS_CRIADO,
                detalhes="OK",
            ),
        )

        resultados = aghu.processar_cadastros(
            context=object(),
            page_inicial=object(),
            janela_sistema_inicial=object(),
            cadastros=[pessoa_valida()],
            usuario_rede="usuario",
            senha="senha",
        )

        assert resultados[0].status == STATUS_CRIADO

    def test_processar_cadastros_retorna_erro_apos_duas_falhas(self, monkeypatch):
        monkeypatch.setattr(
            aghu,
            "garantir_tela_pesquisa_pessoa",
            lambda **_kwargs: ("page", "janela"),
        )
        monkeypatch.setattr(
            aghu,
            "processar_cadastro",
            lambda *_args: (_ for _ in ()).throw(RuntimeError("falha")),
        )
        monkeypatch.setattr(aghu, "trocar_aba_aghux", lambda **_kwargs: "page-limpa")
        monkeypatch.setattr(
            aghu,
            "navegar_ate_cadastro_pessoa",
            lambda **_kwargs: ("page-limpa", "janela-limpa"),
        )

        resultados = aghu.processar_cadastros(
            context=object(),
            page_inicial=object(),
            janela_sistema_inicial=object(),
            cadastros=[pessoa_valida()],
            usuario_rede="usuario",
            senha="senha",
        )

        assert resultados[0].status == STATUS_ERRO
        assert "Falha tecnica definitiva" in resultados[0].detalhes


class FakeBrowser:
    def __init__(self):
        self.context = FakeContext()
        self.fechado = False

    def new_context(self, **kwargs):
        self.context_kwargs = kwargs
        return self.context

    def close(self):
        self.fechado = True


class FakeChromium:
    def __init__(self):
        self.browser = FakeBrowser()
        self.launch_kwargs = None

    def launch(self, **kwargs):
        self.launch_kwargs = kwargs
        return self.browser


class FakePlaywright:
    def __init__(self):
        self.chromium = FakeChromium()


class FakeSyncPlaywright:
    def __init__(self, playwright):
        self.playwright = playwright

    def __enter__(self):
        return self.playwright

    def __exit__(self, exc_type, exc, tb):
        return False


class TestExecucaoPublica:
    @pytest.mark.parametrize(
        ("usuario", "senha", "url", "mensagem"),
        [
            ("", "senha", "url", "usuario de rede"),
            ("usuario", "", "url", "usuario de rede"),
            ("usuario", "senha", "", "ambiente"),
        ],
    )
    def test_executar_cadastro_pessoas_valida_parametros(
        self,
        usuario,
        senha,
        url,
        mensagem,
    ):
        with pytest.raises(ValueError, match=mensagem):
            executar_cadastro_pessoas(
                cadastros=[],
                usuario_rede=usuario,
                senha=senha,
                url_aghu=url,
            )

    def test_executar_cadastro_pessoas_invalidas_nao_abre_playwright(
        self,
        monkeypatch,
        tmp_path,
    ):
        monkeypatch.setattr(
            aghu,
            "sync_playwright",
            lambda: pytest.fail("Playwright nao deveria iniciar"),
        )

        resultados = executar_cadastro_pessoas(
            cadastros=[CadastroPessoaEntrada(cpf="123")],
            usuario_rede="usuario",
            senha="senha",
            diretorio_logs=tmp_path,
        )

        assert resultados[0].status == STATUS_IGNORADO
        assert len(list(tmp_path.glob("log_cadastro_pessoas_*.csv"))) == 1

    def test_executar_cadastro_pessoas_abre_browser_e_propaga_url(
        self,
        monkeypatch,
    ):
        playwright = FakePlaywright()
        resultado = ResultadoCadastroPessoa(
            cpf="12345678901",
            nome_pessoa="Joao",
            status=STATUS_CRIADO,
            detalhes="OK",
        )
        capturados = {}

        monkeypatch.setattr(
            aghu,
            "sync_playwright",
            lambda: FakeSyncPlaywright(playwright),
        )
        monkeypatch.setattr(
            aghu,
            "fazer_login",
            lambda page, usuario_rede, senha, *, url_aghu: capturados.update(
                login=(page, usuario_rede, senha, url_aghu)
            ),
        )
        monkeypatch.setattr(
            aghu,
            "navegar_ate_cadastro_pessoa",
            lambda **kwargs: (
                capturados.update(navegar=kwargs) or kwargs["page_atual"],
                "janela",
            ),
        )
        monkeypatch.setattr(
            aghu,
            "processar_cadastros",
            lambda **kwargs: capturados.update(processar=kwargs) or [resultado],
        )
        monkeypatch.setattr(
            aghu,
            "gerar_csv_logs",
            lambda *_args, **_kwargs: pytest.fail("CSV desabilitado"),
        )

        resultados = executar_cadastro_pessoas(
            cadastros=[pessoa_valida()],
            usuario_rede="usuario",
            senha="senha",
            url_aghu="http://homologacao",
            mostrar_browser=False,
            gerar_csv_log=False,
        )

        browser = playwright.chromium.browser
        assert resultados == [resultado]
        assert playwright.chromium.launch_kwargs["headless"] is True
        assert browser.context_kwargs == {"ignore_https_errors": True}
        assert browser.context.page.goto_url == "http://homologacao"
        assert capturados["login"][3] == "http://homologacao"
        assert capturados["navegar"]["url_aghu"] == "http://homologacao"
        assert capturados["processar"]["url_aghu"] == "http://homologacao"
        assert browser.fechado is True

    def test_executar_cadastro_lote_encadeia_leitura_execucao_e_relatorio(
        self,
        monkeypatch,
    ):
        cadastros = [pessoa_valida()]
        resultados = [
            ResultadoCadastroPessoa(
                cpf="12345678901",
                nome_pessoa="Joao",
                status=STATUS_CRIADO,
                detalhes="OK",
            )
        ]
        capturados = {}
        monkeypatch.setattr(aghu, "ler_planilha_cadastros", lambda caminho: cadastros)
        monkeypatch.setattr(
            aghu,
            "executar_cadastro_pessoas",
            lambda **kwargs: capturados.update(executar=kwargs) or resultados,
        )
        monkeypatch.setattr(
            aghu,
            "salvar_relatorio_resultados",
            lambda resultados_arg, caminho: Path(caminho),
        )

        retorno_resultados, relatorio = executar_cadastro_lote(
            usuario_rede="usuario",
            senha="senha",
            caminho_planilha="entrada.xlsx",
            caminho_relatorio="saida.xlsx",
            url_aghu="url",
            mostrar_browser=False,
            mostrar_console=False,
            diretorio_logs="logs",
        )

        assert retorno_resultados == resultados
        assert relatorio == Path("saida.xlsx")
        assert capturados["executar"]["cadastros"] == cadastros
        assert capturados["executar"]["url_aghu"] == "url"
        assert capturados["executar"]["mostrar_browser"] is False
        assert capturados["executar"]["mostrar_console"] is False
        assert capturados["executar"]["diretorio_logs"] == "logs"

    def test_executar_cadastro_individual_retorna_primeiro_resultado(self, monkeypatch):
        resultado = ResultadoCadastroPessoa(
            cpf="12345678901",
            nome_pessoa="Joao",
            status=STATUS_CRIADO,
            detalhes="OK",
        )
        capturados = {}
        monkeypatch.setattr(
            aghu,
            "executar_cadastro_pessoas",
            lambda **kwargs: capturados.update(kwargs) or [resultado],
        )

        retorno = executar_cadastro_individual(
            usuario_rede="usuario",
            senha="senha",
            cadastro=pessoa_valida(),
            url_aghu="url",
            mostrar_browser=False,
            mostrar_console=False,
            diretorio_logs="logs",
        )

        assert retorno == resultado
        assert len(capturados["cadastros"]) == 1
        assert capturados["url_aghu"] == "url"
        assert capturados["mostrar_browser"] is False
        assert capturados["mostrar_console"] is False
