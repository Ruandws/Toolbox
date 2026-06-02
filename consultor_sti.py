import unicodedata
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
from playwright.sync_api import sync_playwright


LOGIN_URL = "https://servicosti.ebserh.gov.br/#/login"
SEARCH_USERS_URL = "https://servicosti.ebserh.gov.br/#/pesquisa-usuarios"

SUPPORTED_EXTENSIONS = (".xlsx",)
SEARCH_TYPE_CPF = "CPF"
SEARCH_TYPE_FULL_NAME = "Nome Completo"
REPORT_COLUMN_USER = "usuário"
REPORT_COLUMN_FULL_NAME = "Nome Completo"
REPORT_COLUMN_STATUS = "Relatório"

SEARCH_INPUT_PLACEHOLDER = (
    "Informe o e-mail institucional, o CPF, ou o nome do usuário"
)
SEARCH_BUTTON_TEXT = "Pesquisar"
SEARCH_BUTTON_LOADING_TEXT = "Pesquisando..."
SEARCH_BUTTON_SELECTOR = 'button[type="submit"][ng-disabled="pesquisando"]'

NO_USER_FOUND_MESSAGE = "Nenhum usuário encontrado"
NO_USER_FOUND_TEXT = f"{NO_USER_FOUND_MESSAGE}."
NO_USER_FOUND_SELECTOR = 'div[ng-show*="usuarios.length == 0"] .well b'

SEARCH_RESULT_ROW_SELECTOR = (
    'table.table-striped tbody tr[ng-repeat*="usuario in usuarios"]'
)

SEARCH_RESPONSE_TIMEOUT_MS = 5000
SEARCH_LOADING_APPEAR_TIMEOUT_MS = 1500
SEARCH_LOADING_FINISH_TIMEOUT_MS = 10000

LOGIN_BUTTON_TEXT = "Entrar"
LOGIN_PASSWORD_SELECTOR = 'input[type="password"]'
PAGE_READY_TIMEOUT_MS = 15000
LOGIN_COMPLETION_TIMEOUT_MS = 15000

# Filtros Excel
XLSX_HEADER_ROW = 1
XLSX_FREEZE_PANES_CELL = "A2"
XLSX_MIN_COLUMN_WIDTH = 12
XLSX_MAX_COLUMN_WIDTH = 60
XLSX_COLUMN_PADDING = 2

Row = Dict[str, Any]


@dataclass
class SearchResult:
    message: str
    full_name: str = ""
    user_login: str = ""


# -----------------------------
# Planilha / relatório
# -----------------------------

# Normaliza nome de coluna.
def normalize_column_name(value: Any) -> str:    
    text = "" if value is None else str(value).strip()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(
        char for char in text
        if not unicodedata.combining(char)
    )
    return text.lower()


# Identifica coluna alvo da pesquisa.  
def identify_search_column(headers: Sequence[str], search_type: str) -> str:  
    normalized_search_type = normalize_search_type(search_type)
    
    if normalized_search_type == SEARCH_TYPE_CPF:
        candidates = {"cpf"}
    else:
        candidates = {"nome", "nome completo"}

    for header in headers:
        if normalize_column_name(header) in candidates:
            return header

    expected_columns = ", ".join(sorted(candidates))
    raise ValueError(
        "Coluna de pesquisa não encontrada. "
        f"A planilha deve conter uma destas colunas: {expected_columns}."
    )


# Valida extensão da planilha.
def validate_spreadsheet_extension(spreadsheet_path: str) -> str:  
    extension = Path(spreadsheet_path).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(SUPPORTED_EXTENSIONS)
        raise ValueError(
            f"Formato de planilha não suportado: {extension}. "
            f"Use apenas: {supported}."
        )

    return extension


# Gera nome do relatório.
def generate_report_filename(
    extension: str,
    now: Optional[datetime] = None,
) -> str:
    
    reference_date = now or datetime.now()
    clean_extension = extension.lower().strip()

    if not clean_extension.startswith("."):
        clean_extension = f".{clean_extension}"

    timestamp = reference_date.strftime("%d_%m_%y_%Hh%Mm%S")
    return f"Resultado_{timestamp}{clean_extension}"


# Obtém caminho disponível.
def get_available_report_path(report_path: Path) -> Path:
   
    if not report_path.exists():
        return report_path

    counter = 2
    parent = report_path.parent
    stem = report_path.stem
    suffix = report_path.suffix

    while True:
        candidate = parent / f"{stem}_{counter}{suffix}"

        if not candidate.exists():
            return candidate

        counter += 1


# Constrói caminho do relatório.
def build_report_path(
    report_directory: str,
    source_spreadsheet_path: str,
    now: Optional[datetime] = None,
) -> Path:
    
    report_dir = Path(report_directory).expanduser()

    if not report_dir.exists():
        raise FileNotFoundError(
            f"A pasta de relatório não existe: {report_dir}"
        )

    if not report_dir.is_dir():
        raise NotADirectoryError(
            f"O caminho de relatório não é uma pasta: {report_dir}"
        )

    extension = validate_spreadsheet_extension(source_spreadsheet_path)
    report_filename = generate_report_filename(extension, now)
    report_path = report_dir / report_filename

    return get_available_report_path(report_path)


# Lê dados da planilha.
def read_spreadsheet(spreadsheet_path: str) -> Tuple[List[str], List[Row]]:
    
    source_path = Path(spreadsheet_path).expanduser()

    if not source_path.exists():
        raise FileNotFoundError(f"Planilha não encontrada: {source_path}")

    validate_spreadsheet_extension(str(source_path))
    return read_xlsx(source_path)


# Lê arquivo Excel.
def read_xlsx(source_path: Path) -> Tuple[List[str], List[Row]]:
    
    workbook = load_workbook(source_path, data_only=True)
    worksheet = workbook.active

    first_row = next(
        worksheet.iter_rows(min_row=1, max_row=1),
        None,
    )

    if first_row is None:
        raise ValueError("A planilha XLSX está vazia.")

    raw_headers = [cell.value for cell in first_row]
    headers = make_unique_headers(raw_headers)
    rows: List[Row] = []

    for raw_row in worksheet.iter_rows(min_row=2, values_only=True):
        if is_empty_row(raw_row):
            continue

        rows.append(build_row(headers, raw_row))

    return headers, rows


# Garante cabeçalhos únicos.  
def make_unique_headers(raw_headers: Sequence[Any]) -> List[str]:
    
    if not raw_headers:
        raise ValueError("A planilha não possui cabeçalho.")

    headers: List[str] = []
    seen: Dict[str, int] = {}

    for index, raw_header in enumerate(raw_headers, start=1):
        header = "" if raw_header is None else str(raw_header).strip()

        if not header:
            header = f"coluna_{index}"

        if header in seen:
            seen[header] += 1
            header = f"{header}_{seen[header]}"
        else:
            seen[header] = 1

        headers.append(header)

    return headers


# Constrói linha de dados.
def build_row(headers: Sequence[str], raw_row: Sequence[Any]) -> Row:
    row: Row = {}

    for index, header in enumerate(headers):
        row[header] = raw_row[index] if index < len(raw_row) else ""

    return row


# Verifica linha vazia.  
def is_empty_row(raw_row: Sequence[Any]) -> bool:
    return all(
        value is None or str(value).strip() == ""
        for value in raw_row
    )


# Obtém cabeçalhos do relatório.  
def get_report_headers(search_type: str) -> List[str]:
    normalized_search_type = normalize_search_type(search_type)

    if normalized_search_type == SEARCH_TYPE_CPF:
        return [
            REPORT_COLUMN_FULL_NAME,
            REPORT_COLUMN_USER,
            REPORT_COLUMN_STATUS,
        ]

    return [
        REPORT_COLUMN_USER,
        REPORT_COLUMN_STATUS,
    ]


# Constrói linha do relatório.  
def build_report_row(search_type: str, result: SearchResult) -> Row:
    
    report_row: Row = {
        REPORT_COLUMN_USER: result.user_login,
        REPORT_COLUMN_STATUS: result.message,
    }

    if normalize_search_type(search_type) == SEARCH_TYPE_CPF:
        report_row[REPORT_COLUMN_FULL_NAME] = result.full_name

    return report_row


# Salva relatório em arquivo.
def write_report(
    source_spreadsheet_path: str,
    report_path: Path,
    search_type: str,
    rows: Sequence[Row],
) -> None:
    validate_spreadsheet_extension(source_spreadsheet_path)
    report_headers = get_report_headers(search_type)

    write_xlsx_report(report_path, report_headers, rows)


# Salva relatório em Excel.
def write_xlsx_report(
    report_path: Path,
    headers: Sequence[str],
    rows: Sequence[Row],
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Relatório"

    worksheet.append(list(headers))

    for row in rows:
        worksheet.append([
            row.get(header, "")
            for header in headers
        ])

    apply_xlsx_report_layout(worksheet)
    workbook.save(report_path)


# Aplica filtros, congelamento e largura automática no XLSX.
def apply_xlsx_report_layout(worksheet) -> None:
    if worksheet.max_row < XLSX_HEADER_ROW or worksheet.max_column < 1:
        return

    worksheet.freeze_panes = XLSX_FREEZE_PANES_CELL
    worksheet.auto_filter.ref = build_xlsx_filter_range(worksheet)
    autofit_xlsx_columns(worksheet)


# Monta intervalo de filtros do relatório.
def build_xlsx_filter_range(worksheet) -> str:
    last_column = get_column_letter(worksheet.max_column)
    return f"A{XLSX_HEADER_ROW}:{last_column}{worksheet.max_row}"


# Ajusta largura das colunas conforme o maior conteúdo.
def autofit_xlsx_columns(worksheet) -> None:
    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = max(
            get_xlsx_cell_text_length(cell.value)
            for cell in column_cells
        )
        width = max_length + XLSX_COLUMN_PADDING
        worksheet.column_dimensions[column_letter].width = min(
            max(width, XLSX_MIN_COLUMN_WIDTH),
            XLSX_MAX_COLUMN_WIDTH,
        )


# Calcula tamanho visível de uma célula.
def get_xlsx_cell_text_length(value: Any) -> int:
    if value is None:
        return 0

    lines = str(value).splitlines() or [""]
    return max(len(line) for line in lines)


# -----------------------------
# Pesquisa / validação simples
# -----------------------------

# Normaliza tipo de pesquisa.
def normalize_search_type(search_type: str) -> str:
    
    value = str(search_type).strip().lower()

    if value == "cpf":
        return SEARCH_TYPE_CPF

    if value in {"nome completo", "nome", "full name"}:
        return SEARCH_TYPE_FULL_NAME

    raise ValueError(
        "Tipo de pesquisa inválido. Use CPF ou Nome Completo."
    )


# Prepara valor da pesquisa.
def prepare_search_value(search_type: str, search_value: Any) -> str:
    
    normalized_search_type = normalize_search_type(search_type)
    value = "" if search_value is None else str(search_value).strip()

    if not value:
        raise ValueError("Valor da pesquisa não informado.")

    if normalized_search_type == SEARCH_TYPE_CPF:
        value_digits = "".join(char for char in value if char.isdigit())
        if not value_digits:
            raise ValueError(
                "Para pesquisa por CPF, informe um CPF válido."
            )
        return value_digits

    if any(char.isdigit() for char in value):
        raise ValueError(
            "Para pesquisa por Nome Completo, informe o nome sem números."
        )

    return " ".join(value.split())


# -----------------------------
# Autenticação e Navegação
# -----------------------------

# Obtém locator do campo de login.
def get_login_textbox_locator(page):
    return page.get_by_role("textbox").first

# Obtém locator do campo de senha do login.
def get_login_password_locator(page):
    return page.locator(LOGIN_PASSWORD_SELECTOR)


# Aguarda a tela de login carregar após DOMContentLoaded.
def wait_for_login_page_ready(page) -> None:
    get_login_textbox_locator(page).wait_for(
        state="visible",
        timeout=PAGE_READY_TIMEOUT_MS,
    )
    get_login_password_locator(page).wait_for(
        state="visible",
        timeout=PAGE_READY_TIMEOUT_MS,
    )

# Aguarda o login concluir sem usar espera fixa.
def wait_for_login_completion(page) -> None:
    try:
        get_login_password_locator(page).wait_for(
            state="hidden",
            timeout=LOGIN_COMPLETION_TIMEOUT_MS,
        )
    except PlaywrightTimeoutError as exc:
        raise RuntimeError(
            "Login não concluído no tempo esperado. "
            "Verifique as credenciais, a disponibilidade do sistema "
            "ou mudança no fluxo de autenticação."
        ) from exc


# Realiza login no sistema.
def login_to_system(page, login: str, password: str) -> None:
    
    page.goto(LOGIN_URL, wait_until="domcontentloaded")
    wait_for_login_page_ready(page)
    get_login_textbox_locator(page).fill(login)
    get_login_password_locator(page).fill(password)
    page.get_by_role("button", name=LOGIN_BUTTON_TEXT).click()
    wait_for_login_completion(page)


# Abre tela de pesquisa.
def open_search_users_page(page) -> None:

    page.goto(SEARCH_USERS_URL, wait_until="domcontentloaded")
    wait_for_search_users_page_ready(page)


# -----------------------------
# Locators e Esperas de Pesquisa
# -----------------------------

# Aguarda a tela de pesquisa ficar pronta.
def wait_for_search_users_page_ready(page) -> None:
    get_search_input_locator(page).wait_for(
        state="visible",
        timeout=PAGE_READY_TIMEOUT_MS,
    )

# Obtém locator do campo de pesquisa.
def get_search_input_locator(page):
    return page.get_by_placeholder(SEARCH_INPUT_PLACEHOLDER)


# Obtém locator do botão de pesquisa.
def get_search_button_locator(page):
    return page.locator(SEARCH_BUTTON_SELECTOR).filter(
        has_text=SEARCH_BUTTON_TEXT,
    )


# Obtém locator do botão em estado de carregamento.
def get_search_loading_button_locator(page):
    return page.locator(SEARCH_BUTTON_SELECTOR).filter(
        has_text=SEARCH_BUTTON_LOADING_TEXT,
    )


# Obtém locator da mensagem explícita de ausência de resultado.
def get_no_user_found_locator(page):
    return page.locator(NO_USER_FOUND_SELECTOR).filter(
        has_text=NO_USER_FOUND_TEXT,
    )


# Obtém locator das linhas válidas de usuário.
def get_search_result_rows_locator(page):
    return page.locator(SEARCH_RESULT_ROW_SELECTOR)


# Aguarda o ciclo de carregamento da pesquisa.
def wait_for_search_loading_cycle(page) -> None:
    loading_button = get_search_loading_button_locator(page)

    try:
        loading_button.wait_for(
            state="visible",
            timeout=SEARCH_LOADING_APPEAR_TIMEOUT_MS,
        )
    except PlaywrightTimeoutError:
        return

    loading_button.wait_for(
        state="hidden",
        timeout=SEARCH_LOADING_FINISH_TIMEOUT_MS,
    )


# Aguarda a resposta da pesquisa: linha de resultado ou mensagem de ausência.
def wait_for_search_response(page) -> None:
    result_row = get_search_result_rows_locator(page).first
    no_user_found = get_no_user_found_locator(page).first

    result_row.or_(no_user_found).first.wait_for(
        state="visible",
        timeout=SEARCH_RESPONSE_TIMEOUT_MS,
    )

# Extrai resultado da tabela.
def extract_single_result(row, search_type: str) -> SearchResult:

    normalized_search_type = normalize_search_type(search_type)
    cells = row.locator("td")
    full_name = cells.nth(3).inner_text().strip()
    user_login = cells.nth(4).inner_text().strip()

    if normalized_search_type == SEARCH_TYPE_CPF:
        return SearchResult(
            message="Usuário encontrado",
            full_name=full_name,
            user_login=user_login,
        )

    return SearchResult(
        message="Usuário encontrado",
        user_login=user_login,
    )


# Pesquisa o usuário.
def search_user(page, search_type: str, search_value: Any) -> SearchResult:
    
    clean_value = prepare_search_value(search_type, search_value)

    input_campo = get_search_input_locator(page)
    input_campo.wait_for(state="visible")
    input_campo.fill("")
    input_campo.fill(clean_value)

    get_search_button_locator(page).click()

    try:
        wait_for_search_loading_cycle(page)
        wait_for_search_response(page)

        if get_no_user_found_locator(page).first.is_visible():
            return SearchResult(message=NO_USER_FOUND_MESSAGE)

        rows = get_search_result_rows_locator(page)
        count = rows.count()

        if count == 0:
            return SearchResult(message=NO_USER_FOUND_MESSAGE)

        if count > 1:
            return SearchResult(message="Mais de um usuário encontrado")

        return extract_single_result(rows.first, search_type)
    except PlaywrightTimeoutError:
        return SearchResult(message=NO_USER_FOUND_MESSAGE)


# Formata resultado da pesquisa.
def format_single_result(result: SearchResult, search_type: str) -> str:
    
    if result.message != "Usuário encontrado":
        return result.message

    if normalize_search_type(search_type) == SEARCH_TYPE_CPF:
        return (
            "Usuário encontrado!\n"
            f"Nome Completo: {result.full_name}\n"
            f"Usuário: {result.user_login}"
        )

    return (
        "Usuário encontrado!\n"
        f"Usuário: {result.user_login}"
    )


# -----------------------------
# Entry points
# -----------------------------

# Executa automação individual.
def run_automation(
    login: str,
    password: str,
    search_type: str,
    search_value: str,
) -> str:
    
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=False)
        context = browser.new_context()

        try:
            page = context.new_page()
            login_to_system(page, login, password)
            open_search_users_page(page)
            result = search_user(page, search_type, search_value)
            return format_single_result(result, search_type)
        finally:
            context.close()
            browser.close()


# Executa automação em lote.
def run_batch_automation(
    login: str,
    password: str,
    search_type: str,
    spreadsheet_path: str,
    report_directory: str,
) -> str:
    
    normalized_search_type = normalize_search_type(search_type)
    headers, source_rows = read_spreadsheet(spreadsheet_path)
    search_column = identify_search_column(headers, normalized_search_type)
    report_path = build_report_path(report_directory, spreadsheet_path)
    report_rows: List[Row] = []

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=False)
        context = browser.new_context()

        try:
            page = context.new_page()
            login_to_system(page, login, password)
            open_search_users_page(page)

            for source_row in source_rows:
                search_value = source_row.get(search_column, "")

                try:
                    result = search_user(
                        page,
                        normalized_search_type,
                        search_value,
                    )
                except ValueError as exc:
                    result = SearchResult(message=f"Erro: {str(exc)}")
                except Exception as exc:
                    result = SearchResult(message=f"Erro: {str(exc)}")

                report_rows.append(
                    build_report_row(normalized_search_type, result)
                )
        finally:
            context.close()
            browser.close()

    write_report(
        spreadsheet_path,
        report_path,
        normalized_search_type,
        report_rows,
    )

    return (
        f"Lote finalizado. {len(report_rows)} registro(s) processado(s). "
        f"Relatório: {report_path}"
    )
